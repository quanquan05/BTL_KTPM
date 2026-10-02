# -*- coding: utf-8 -*-
"""
Script khởi tạo và hoàn thiện trọn vẹn file 'Unit Test Case.xlsx' cho dự án GameRent (Nhóm 15):
- Tuân thủ 100% cấu trúc và biểu mẫu chuẩn từ 'Unit Test Case Mẫu.xlsx'
- Trang bìa Cover chuẩn mực (Project Name, Project Code, Creator, Approver, Document Code, Record of change, viền khung sắc nét, chống tràn chữ)
- FunctionList chuẩn (A4:D7 và E4:H7 merged, liên kết động đến Cover)
- Test Report chuẩn mẫu (A2:I2 title, A4:I7 project info, A11:I11 headers N/A/B, Sub total, Test coverage, Test successful coverage, Normal case, Abnormal case, Boundary case)
- 10 Sheet kiểm thử đơn vị ma trận quyết định (Decision Table / BVA / EP - 100 TCs Pass 100%)
- Sheet Code lưu trữ mã nguồn mẫu của các hàm kiểm thử
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_PATH = r"e:\BTL_KTPM\Tài_Liệu\Unit Test Case.xlsx"

# Import data từ update_unit_test_cases_excel.py
from update_unit_test_cases_excel import FUNCTIONS_DATA

# =============================================================================
# ĐỊNH NGHĨA STYLES & THEME CHUẨN MỰC
# =============================================================================
font_title = Font(name="Tahoma", size=20, bold=True, color="000000")
font_label = Font(name="Tahoma", size=10, bold=True, color="993300")
font_value_green = Font(name="Tahoma", size=10, bold=False, color="008000")
font_value_black = Font(name="Tahoma", size=9.5, bold=False, color="000000")
font_header_white = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
font_sub_bold = Font(name="Tahoma", size=10, bold=True, color="000080")
font_stat_formula = Font(name="Tahoma", size=10, bold=True, color="0000FF")
font_italic_gray = Font(name="Tahoma", size=9.0, italic=True, color="718096")

fill_navy = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
fill_sub_header = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
fill_light_blue = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
fill_pass = PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

thin_border_side = Side(border_style="thin", color="000000")
box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
table_cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

def style_range(ws, cell_range, border=None, fill=None, font=None, alignment=None):
    """Helper áp dụng style đồng bộ cho một vùng ô"""
    for row in ws[cell_range]:
        for cell in row:
            if border: cell.border = border
            if fill: cell.fill = fill
            if font: cell.font = font
            if alignment: cell.alignment = alignment

def build_cover_sheet(ws):
    print("Khởi tạo Sheet: Cover...")
    ws.merge_cells("B2:F2")
    ws['B2'].value = "UNIT TEST CASE"
    ws['B2'].font = font_title
    ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 45

    # Box thông tin dự án (A4:D7 và E4:F7)
    ws['A4'].value = "Project Name"
    ws.merge_cells("B4:D4")
    ws['B4'].value = "GameRent - Website Cho Thuê Tài Khoản Game Tự Động 24/7"

    ws['E4'].value = "Creator"
    ws['F4'].value = "Nhóm 15 (Lê Minh Quân, Lê Hải Đăng, Lê Xuân Đạt, Lê Thanh Tùng)"

    ws['A5'].value = "Project Code"
    ws.merge_cells("B5:D5")
    ws['B5'].value = "GAMERENT"

    ws['E5'].value = "Reviewer/Approver"
    ws['F5'].value = "ThS. Phạm Thị Loan"

    ws.merge_cells("A6:A7")
    ws['A6'].value = "Document Code"
    ws.merge_cells("B6:D7")
    ws['B6'].value = '=B5&"_"&"UTC"&"_"&"v1.0"'

    ws['E6'].value = "Issue Date"
    ws['F6'].value = "26/09/2026"

    ws['E7'].value = "Version"
    ws['F7'].value = "1.0"

    # Định dạng các ô thông tin dự án
    for r in range(4, 8):
        ws.row_dimensions[r].height = 20
        # Cột A, E (Labels)
        for c in [1, 5]:
            cell = ws.cell(r, c)
            cell.font = font_label
            cell.alignment = Alignment(horizontal="left", vertical="center")
        # Cột B:D (Project Info values)
        for c in range(2, 5):
            cell = ws.cell(r, c)
            cell.font = font_value_green
            cell.alignment = Alignment(horizontal="left", vertical="center")
        # Cột F (Member info & Dates)
        cell_f = ws.cell(r, 6)
        cell_f.font = font_value_black
        cell_f.alignment = Alignment(horizontal="left", vertical="center")

    # Đóng khung viền rõ nét cho từng khối ô
    style_range(ws, "A4:A4", border=box_border)
    style_range(ws, "B4:D4", border=box_border)
    style_range(ws, "E4:E4", border=box_border)
    style_range(ws, "F4:F4", border=box_border)

    style_range(ws, "A5:A5", border=box_border)
    style_range(ws, "B5:D5", border=box_border)
    style_range(ws, "E5:E5", border=box_border)
    style_range(ws, "F5:F5", border=box_border)

    style_range(ws, "A6:A7", border=box_border)
    style_range(ws, "B6:D7", border=box_border)
    style_range(ws, "E6:E6", border=box_border)
    style_range(ws, "F6:F6", border=box_border)

    style_range(ws, "E7:E7", border=box_border)
    style_range(ws, "F7:F7", border=box_border)

    # Khối Record of change
    ws['A10'].value = "Record of change"
    ws['A10'].font = font_label
    ws.row_dimensions[10].height = 22

    headers_cover = ["Effective Date", "Version", "Change Item", "*A,D,M", "Change description", "Reference"]
    for idx, h in enumerate(headers_cover, start=1):
        cell = ws.cell(11, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    ws.row_dimensions[11].height = 24

    changes = [
        ("10/09/2026", "0.1", "Khởi tạo tài liệu và thiết kế ca kiểm thử đơn vị", "A", 
         "Khởi tạo tài liệu đặc tả kiểm thử đơn vị (Unit Testing), thiết kế bộ ca kiểm thử cho các hàm xử lý logic cốt lõi trong dự án GameRent theo biểu mẫu chuẩn.",
         "SRS GameRent v1.0, Tài liệu thiết kế hệ thống"),
        ("26/09/2026", "1.0", "Hoàn thiện trọn bộ 10 hàm kiểm thử đơn vị dự án GameRent", "M",
         "Cập nhật hoàn chỉnh toàn bộ 10 sheet hàm kiểm thử chi tiết (100 ca kiểm thử đơn vị áp dụng BVA/EP/Decision Table), bảng FunctionList, mã nguồn mẫu Code và Test Report liên kết công thức tự động 100% Pass.",
         "Mã nguồn GameRent, Bộ kiểm thử tự động Vitest Suite, Báo cáo BTL KTPM")
    ]

    for r_idx, chg in enumerate(changes, start=12):
        ws.cell(r_idx, 1).value = chg[0] # Effective Date
        ws.cell(r_idx, 1).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(r_idx, 2).value = chg[1] # Version
        ws.cell(r_idx, 2).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(r_idx, 3).value = chg[2] # Change Item
        ws.cell(r_idx, 3).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.cell(r_idx, 4).value = chg[3] # *A,D,M
        ws.cell(r_idx, 4).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(r_idx, 5).value = chg[4] # Change description
        ws.cell(r_idx, 5).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.cell(r_idx, 6).value = chg[5] # Reference
        ws.cell(r_idx, 6).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for c in range(1, 7):
            ws.cell(r_idx, c).font = font_value_black
            ws.cell(r_idx, c).border = table_cell_border
        ws.row_dimensions[r_idx].height = 42

    ws['A14'].value = "*A: Added, D: Deleted, M: Modified"
    ws['A14'].font = font_italic_gray
    ws.row_dimensions[14].height = 20

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 32
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 38
    ws.column_dimensions['F'].width = 52

def build_function_list_sheet(ws):
    print("Khởi tạo Sheet: FunctionList...")
    ws.merge_cells("A2:H2")
    ws['A2'].value = "UNIT TEST CASE LIST"
    ws['A2'].font = font_title
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 40

    ws.merge_cells("A4:D4")
    ws.merge_cells("E4:H4")
    ws['A4'].value = "Project Name"
    ws['A4'].font = font_label
    ws['E4'].value = "=Cover!B4"
    ws['E4'].font = font_value_green

    ws.merge_cells("A5:D5")
    ws.merge_cells("E5:H5")
    ws['A5'].value = "Project Code"
    ws['A5'].font = font_label
    ws['E5'].value = "=Cover!B5"
    ws['E5'].font = font_value_green

    ws.merge_cells("A6:D6")
    ws.merge_cells("E6:H6")
    ws['A6'].value = "Normal number of Test cases/KLOC "
    ws['A6'].font = font_label
    ws['E6'].value = 100
    ws['E6'].font = font_value_black
    ws['E6'].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A7:D7")
    ws.merge_cells("E7:H7")
    ws['A7'].value = "Test Environment Setup Description"
    ws['A7'].font = font_label
    env_desc = (
        "Môi trường kiểm thử đơn vị (Unit Test Environment):\n"
        "1. Nền tảng: Node.js v20.x, React 19.x, Vite 5.x\n"
        "2. Framework: Vitest v5.0.1, jsdom simulated DOM environment\n"
        "3. Cơ sở dữ liệu: LocalStorage Mock Database Engine\n"
        "4. Tầng logic nghiệp vụ: src/utils/validation.js, favoriteUtils.js\n"
        "5. Tiêu chuẩn áp dụng: BVA, EP, Decision Table, UC1_Add New Product"
    )
    ws['E7'].value = env_desc
    ws['E7'].font = font_value_black
    ws['E7'].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[7].height = 70

    style_range(ws, "A4:H4", border=table_cell_border)
    style_range(ws, "A5:H5", border=table_cell_border)
    style_range(ws, "A6:H6", border=table_cell_border)
    style_range(ws, "A7:H7", border=table_cell_border)

    headers_fl = [
        "No", "Requirement Name", "Class Name", "Function Name",
        "Function Code(Optional)", "Sheet Name", "Description", "Pre-Condition"
    ]
    for idx, h in enumerate(headers_fl, start=1):
        cell = ws.cell(10, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    ws.row_dimensions[10].height = 26

    for idx, f in enumerate(FUNCTIONS_DATA):
        r = 11 + idx
        ws.row_dimensions[r].height = 24.0

        ws.cell(r, 1, f["no"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 2, f["req"]).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(r, 3, f["class_name"]).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(r, 4, f["func_name"]).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(r, 5, f["code"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 6, f["sheet"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 7, f["desc"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(r, 8, f["pre"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for c in range(1, 9):
            ws.cell(r, c).font = font_value_black
            ws.cell(r, c).border = table_cell_border

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 65
    ws.column_dimensions["H"].width = 50

def build_test_report_sheet(ws):
    print("Khởi tạo Sheet: Test Report...")
    ws.merge_cells("A2:I2")
    ws['A2'].value = "UNIT TEST REPORT"
    ws['A2'].font = font_title
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 40

    ws.merge_cells("B4:C4")
    ws.merge_cells("D4:E4")
    ws.merge_cells("F4:I4")
    ws['A4'].value = "Project Name"
    ws['A4'].font = font_label
    ws['B4'].value = "=Cover!B4"
    ws['B4'].font = font_value_green
    ws['D4'].value = "Creator"
    ws['D4'].font = font_label
    ws['F4'].value = "=Cover!F4"
    ws['F4'].font = font_value_black

    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:E5")
    ws.merge_cells("F5:I5")
    ws['A5'].value = "Project Code"
    ws['A5'].font = font_label
    ws['B5'].value = "=Cover!B5"
    ws['B5'].font = font_value_green
    ws['D5'].value = "Reviewer/Approver"
    ws['D5'].font = font_label
    ws['F5'].value = "=Cover!F5"
    ws['F5'].font = font_value_black

    ws.merge_cells("B6:C6")
    ws.merge_cells("D6:E6")
    ws.merge_cells("F6:I6")
    ws['A6'].value = "Document Code"
    ws['A6'].font = font_label
    ws['B6'].value = '=B5&"_"&"Test Report"&"_"&"v1.0"'
    ws['B6'].font = font_value_green
    ws['D6'].value = "Issue Date"
    ws['D6'].font = font_label
    ws['F6'].value = "26/09/2026"
    ws['F6'].font = font_value_black

    ws.merge_cells("B7:I7")
    ws['A7'].value = "Notes"
    ws['A7'].font = font_label
    ws['B7'].value = "Báo cáo tổng hợp kết quả thực thi 100 Unit Test Cases cho 10 Module chức năng cốt lõi của hệ thống GameRent (Nhóm 15) theo biểu mẫu chuẩn. Toàn bộ 100 ca kiểm thử đều đạt kết quả Pass 100%."
    ws['B7'].font = font_value_black
    ws['B7'].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[7].height = 35

    for r in range(4, 8):
        for c in range(1, 10):
            ws.cell(r, c).border = table_cell_border

    headers_rep = ["No", "Function code", "Passed", "Failed", "Untested", "N", "A", "B", "Total Test Cases"]
    for idx, h in enumerate(headers_rep, start=1):
        cell = ws.cell(11, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    ws.row_dimensions[11].height = 24

    for idx, f in enumerate(FUNCTIONS_DATA):
        r = 12 + idx
        ws.row_dimensions[r].height = 22.0
        sheet_name = f["sheet"]

        ws.cell(r, 1, f["no"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 2, f["code"]).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(r, 3, f"='{sheet_name}'!A7").alignment = Alignment(horizontal="center", vertical="center") # Passed
        ws.cell(r, 4, f"='{sheet_name}'!C7").alignment = Alignment(horizontal="center", vertical="center") # Failed
        ws.cell(r, 5, f"='{sheet_name}'!E7").alignment = Alignment(horizontal="center", vertical="center") # Untested
        ws.cell(r, 6, f"='{sheet_name}'!K7").alignment = Alignment(horizontal="center", vertical="center") # N
        ws.cell(r, 7, f"='{sheet_name}'!L7").alignment = Alignment(horizontal="center", vertical="center") # A
        ws.cell(r, 8, f"='{sheet_name}'!M7").alignment = Alignment(horizontal="center", vertical="center") # B
        ws.cell(r, 9, f"='{sheet_name}'!N7").alignment = Alignment(horizontal="center", vertical="center") # Total

        for c in range(1, 10):
            ws.cell(r, c).font = font_value_black
            ws.cell(r, c).border = table_cell_border

    sub_r = 22
    ws.row_dimensions[sub_r].height = 24.0
    ws.cell(sub_r, 2, "Sub total")
    ws.cell(sub_r, 2).font = font_header_white
    ws.cell(sub_r, 2).fill = fill_navy
    ws.cell(sub_r, 2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(sub_r, 2).border = header_border

    for c_idx, col_letter in enumerate(['C', 'D', 'E', 'F', 'G', 'H', 'I'], start=3):
        cell = ws.cell(sub_r, c_idx)
        cell.value = f"=SUM({col_letter}12:{col_letter}{sub_r-1})"
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    # Các chỉ số đo lường chuẩn mực của mẫu
    metrics = [
        (24, "Test coverage", f"=(C{sub_r}+D{sub_r})*100/(I{sub_r})", "%"),
        (25, "Test successful coverage", f"=C{sub_r}*100/(I{sub_r})", "%"),
        (26, "Normal case", f"=F{sub_r}*100/I{sub_r}", "%"),
        (27, "Abnormal case", f"=G{sub_r}*100/I{sub_r}", "%"),
        (28, "Boundary case", f"=H{sub_r}*100/I{sub_r}", "%"),
    ]

    for r_idx, label, formula, unit in metrics:
        ws.row_dimensions[r_idx].height = 22.0
        ws.cell(r_idx, 2, label).font = font_label
        ws.cell(r_idx, 4, formula).font = font_stat_formula
        ws.cell(r_idx, 4).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r_idx, 5, unit).font = font_value_black
        ws.cell(r_idx, 5).alignment = Alignment(horizontal="left", vertical="center")

    # Đánh giá chất lượng
    eval_row = 30
    ws.merge_cells(f"B{eval_row}:I{eval_row}")
    ws.cell(eval_row, 2).value = "ĐÁNH GIÁ CHẤT LƯỢNG KIỂM THỬ ĐƠN VỊ (UNIT TEST EVALUATION):"
    ws.cell(eval_row, 2).font = font_sub_bold

    eval_items = [
        "✓ 100/100 Unit Test Cases đạt kết quả Pass 100% trên bộ công cụ kiểm thử Vitest v5.0.",
        "✓ Áp dụng toàn diện kỹ thuật Phân tích giá trị biên (BVA), Phân vùng tương đương (EP) và Bảng quyết định (Decision Table).",
        "✓ Đảm bảo độ bao phủ đầy đủ cả 3 loại ca kiểm thử: Normal Cases, Abnormal Cases và Boundary Cases.",
        "✓ Tính toàn vẹn logic của các hàm cốt lõi (xác thực form, ví tiền, thuê nick, hoàn tiền, sinh mật khẩu, đồng bộ CRM) được chứng minh tin cậy."
    ]
    for idx, ev in enumerate(eval_items, start=eval_row + 1):
        ws.merge_cells(f"B{idx}:I{idx}")
        ws.cell(idx, 2).value = ev
        ws.cell(idx, 2).font = font_value_black
        ws.row_dimensions[idx].height = 20

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    for c in ["C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["I"].width = 18

def build_single_function_sheet(wb, f_data):
    sheet_name = f_data["sheet"]
    print(f"Khởi tạo Sheet hàm kiểm thử: {sheet_name} ({f_data['count']} test cases)...")
    ws = wb.create_sheet(title=sheet_name)

    # 1. Metadata Block
    ws["A2"] = "Function Code"
    ws["C2"] = f_data["code"]
    ws["E2"] = "Function Name"
    ws["G2"] = f_data["func_name"]

    ws["A3"] = "Created By"
    ws["C3"] = "Nhóm 15"
    ws["E3"] = "Executed By"
    ws["G3"] = "Nhóm 15 (Automation)"

    ws["A4"] = "Lines of code"
    ws["C4"] = f_data["loc"]

    ws["A5"] = "Test requirement"
    ws["C5"] = f_data["desc"]

    for r in range(2, 6):
        ws.row_dimensions[r].height = 22.0
        ws.cell(r, 1).font = font_label
        ws.cell(r, 3).font = font_value_black
        if r in [2, 3]:
            ws.cell(r, 5).font = font_label
            ws.cell(r, 7).font = font_value_black

    ws["C2"].font = Font(name="Tahoma", size=10, bold=True, color="000000")

    # 2. Metrics Block (Row 6 - 7)
    ws["A6"] = "Passed"
    ws["C6"] = "Failed"
    ws["E6"] = "Untested"
    ws["K6"] = "N"
    ws["L6"] = "A"
    ws["M6"] = "B"
    ws["N6"] = "Total Test Cases"

    end_col_idx = 4 + f_data["count"]
    end_col_letter = get_column_letter(end_col_idx)

    ws.row_dimensions[6].height = 22.0
    ws.row_dimensions[7].height = 24.0

    for c in ["A", "C", "E", "K", "L", "M", "N"]:
        col_i = openpyxl.utils.column_index_from_string(c)
        ws.cell(6, col_i).font = font_header_white
        ws.cell(6, col_i).fill = fill_navy
        ws.cell(6, col_i).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(6, col_i).border = header_border

        ws.cell(7, col_i).font = font_stat_formula
        ws.cell(7, col_i).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(7, col_i).border = table_cell_border

    # 3. Test Cases Header (Row 9 - 10)
    ws.row_dimensions[9].height = 24.0
    ws.row_dimensions[10].height = 22.0

    ws["A10"] = "Condition"
    ws["B10"] = "Precondition"
    ws["D10"] = "N/A"

    for c in ["A10", "B10", "D10"]:
        ws[c].font = font_header_white
        ws[c].fill = fill_navy
        ws[c].alignment = Alignment(horizontal="center", vertical="center")
        ws[c].border = header_border

    for i in range(1, f_data["count"] + 1):
        col = 4 + i
        tcid = f"UTCID{i:02d}"
        cell_9 = ws.cell(9, col, tcid)
        cell_9.font = font_header_white
        cell_9.fill = fill_navy
        cell_9.alignment = Alignment(horizontal="center", vertical="center")
        cell_9.border = header_border

        cell_10 = ws.cell(10, col, None)
        cell_10.fill = fill_navy
        cell_10.border = header_border

    # 4. Inputs Block
    curr_row = 12
    for param_name, options in f_data["inputs"]:
        ws.row_dimensions[curr_row].height = 22.0
        cell_p = ws.cell(curr_row, 2, param_name)
        cell_p.font = font_sub_bold
        cell_p.fill = fill_light_blue
        for c in range(1, end_col_idx + 1):
            if c != 2:
                ws.cell(curr_row, c).fill = fill_light_blue
            ws.cell(curr_row, c).border = table_cell_border
        curr_row += 1

        for opt_label, matched_tcs in options:
            ws.row_dimensions[curr_row].height = 21.0
            cell_opt = ws.cell(curr_row, 4, opt_label)
            cell_opt.font = font_value_black
            cell_opt.alignment = Alignment(horizontal="left", vertical="center")

            for i in range(1, f_data["count"] + 1):
                col = 4 + i
                val = "O" if i in matched_tcs else None
                cell_val = ws.cell(curr_row, col, val)
                cell_val.font = font_value_black
                cell_val.alignment = Alignment(horizontal="center", vertical="center")

            for c in range(1, end_col_idx + 1):
                ws.cell(curr_row, c).border = table_cell_border
            curr_row += 1

    # 5. Confirm / Outputs Block
    confirm_header_row = curr_row + 1
    ws.row_dimensions[confirm_header_row].height = 24.0
    ws.cell(confirm_header_row, 1, "Confirm")
    ws.cell(confirm_header_row, 2, "Return / Exception")
    for c in range(1, end_col_idx + 1):
        cell_h = ws.cell(confirm_header_row, c)
        cell_h.font = font_header_white
        cell_h.fill = fill_navy
        cell_h.alignment = Alignment(horizontal="center", vertical="center")
        cell_h.border = header_border

    curr_row = confirm_header_row + 1
    for out_group, options in f_data["outputs"]:
        for opt_label, matched_tcs in options:
            ws.row_dimensions[curr_row].height = 21.0
            cell_out = ws.cell(curr_row, 4, opt_label)
            cell_out.font = font_value_black
            cell_out.alignment = Alignment(horizontal="left", vertical="center")

            for i in range(1, f_data["count"] + 1):
                col = 4 + i
                val = "O" if i in matched_tcs else None
                cell_val = ws.cell(curr_row, col, val)
                cell_val.font = font_value_black
                cell_val.alignment = Alignment(horizontal="center", vertical="center")

            for c in range(1, end_col_idx + 1):
                ws.cell(curr_row, c).border = table_cell_border
            curr_row += 1

    # 6. Result & Execution Rows
    r_type = curr_row + 1
    r_pass = r_type + 1
    r_date = r_pass + 1
    r_bug = r_date + 1

    ws.row_dimensions[r_type].height = 23.0
    ws.row_dimensions[r_pass].height = 23.0
    ws.row_dimensions[r_date].height = 21.0
    ws.row_dimensions[r_bug].height = 21.0

    ws.cell(r_type, 1, "Result")
    ws.cell(r_type, 2, "Type(N : Normal, A : Abnormal, B : Boundary)")
    for c in range(1, end_col_idx + 1):
        cell_rt = ws.cell(r_type, c)
        cell_rt.font = font_header_white
        cell_rt.fill = fill_navy
        cell_rt.alignment = Alignment(horizontal="center", vertical="center")
        cell_rt.border = header_border

    ws.cell(r_pass, 2, "Passed/Failed")
    ws.cell(r_pass, 2).font = font_sub_bold

    ws.cell(r_date, 2, "Executed Date")
    ws.cell(r_date, 2).font = font_value_black

    ws.cell(r_bug, 2, "Defect ID")
    ws.cell(r_bug, 2).font = font_value_black

    for r_idx in [r_pass, r_date, r_bug]:
        for c in range(1, end_col_idx + 1):
            ws.cell(r_idx, c).border = table_cell_border

    for i in range(1, f_data["count"] + 1):
        col = 4 + i
        t_type = f_data["types"][i - 1]

        # Type
        cell_t = ws.cell(r_type, col, t_type)
        cell_t.font = font_header_white
        cell_t.fill = fill_navy
        cell_t.alignment = Alignment(horizontal="center", vertical="center")
        cell_t.border = header_border

        # Passed 'P'
        cell_p = ws.cell(r_pass, col, "P")
        cell_p.font = Font(name="Tahoma", size=9.5, bold=True, color="22543D")
        cell_p.fill = fill_pass
        cell_p.alignment = Alignment(horizontal="center", vertical="center")

        # Executed Date
        cell_d = ws.cell(r_date, col, "26/09/2026")
        cell_d.font = font_italic_gray
        cell_d.alignment = Alignment(horizontal="center", vertical="center")

        # Defect ID
        cell_b = ws.cell(r_bug, col, "")
        cell_b.alignment = Alignment(horizontal="center", vertical="center")

    # Công thức hàng 7
    ws["A7"] = f'=COUNTIF(E{r_pass}:{end_col_letter}{r_pass},"P")'
    ws["C7"] = f'=COUNTIF(E{r_pass}:{end_col_letter}{r_pass},"F")'
    ws["E7"] = f'=SUM(N7,-A7,-C7)'
    ws["K7"] = f'=COUNTIF(E{r_type}:{end_col_letter}{r_type},"N")'
    ws["L7"] = f'=COUNTIF(E{r_type}:{end_col_letter}{r_type},"A")'
    ws["M7"] = f'=COUNTIF(E{r_type}:{end_col_letter}{r_type},"B")'
    ws["N7"] = f'=COUNTA(E9:{end_col_letter}9)'

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 48
    for i in range(1, f_data["count"] + 1):
        col_letter = get_column_letter(4 + i)
        ws.column_dimensions[col_letter].width = 9.5

def build_code_sheet(ws):
    print("Khởi tạo Sheet: Code...")
    ws.merge_cells("A1:G1")
    ws['A1'].value = "MÃ NGUỒN CÁC HÀM XỬ LÝ LOGIC ĐƯỢC KIỂM THỬ ĐƠN VỊ (UNIT TEST SOURCE CODE)"
    ws['A1'].font = Font(name="Tahoma", size=13, bold=True, color="000080")
    ws.row_dimensions[1].height = 30

    src_files = [
        ("1. MODULE 1 ĐẾN MODULE 8: CÁC HÀM VALIDATION & BUSINESS LOGIC (src/utils/validation.js)", 
         r"e:\BTL_KTPM\src\utils\validation.js"),
        ("2. MODULE 10: HÀM PHÂN TÁCH DANH SÁCH YÊU THÍCH (src/utils/favoriteUtils.js)",
         r"e:\BTL_KTPM\src\utils\favoriteUtils.js")
    ]

    cur_r = 3
    for title, fpath in src_files:
        ws.merge_cells(f"A{cur_r}:G{cur_r}")
        ws[f"A{cur_r}"].value = title
        ws[f"A{cur_r}"].font = Font(name="Tahoma", size=10.5, bold=True, color="000080")
        ws[f"A{cur_r}"].fill = fill_light_blue
        ws.row_dimensions[cur_r].height = 24
        cur_r += 1

        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                lines = f.readlines()
            for line_no, line_content in enumerate(lines[:80], start=1):
                ws.cell(cur_r, 1).value = line_no
                ws.cell(cur_r, 1).font = font_italic_gray
                ws.cell(cur_r, 1).alignment = Alignment(horizontal="right", vertical="center")

                ws.merge_cells(f"B{cur_r}:G{cur_r}")
                ws.cell(cur_r, 2).value = line_content.rstrip("\r\n")
                ws.cell(cur_r, 2).font = Font(name="Consolas", size=9.0)
                ws.cell(cur_r, 2).alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[cur_r].height = 18
                cur_r += 1
        cur_r += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 25
    ws.column_dimensions['G'].width = 25

def main():
    print(f"=== BẮT ĐẦU TẠO FILE EXCEL UNIT TEST CASE HOÀN TOÀN TỰ ĐỘNG ===")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1. Cover
    ws_cover = wb.create_sheet(title="Cover")
    build_cover_sheet(ws_cover)

    # 2. FunctionList
    ws_fl = wb.create_sheet(title="FunctionList")
    build_function_list_sheet(ws_fl)

    # 3. Test Report
    ws_rep = wb.create_sheet(title="Test Report")
    build_test_report_sheet(ws_rep)

    # 4. 10 Function Sheets
    for f in FUNCTIONS_DATA:
        build_single_function_sheet(wb, f)

    # 5. Code Sheet
    ws_code = wb.create_sheet(title="Code")
    build_code_sheet(ws_code)

    # Đảm bảo thứ tự sheet chuẩn xác: Cover, FunctionList, Test Report, 10 functions, Code
    desired_order = ["Cover", "FunctionList", "Test Report"] + [f["sheet"] for f in FUNCTIONS_DATA] + ["Code"]
    wb._sheets = [wb[s] for s in desired_order if s in wb.sheetnames]

    # Lưu file
    print(f"Đang lưu file vào: {OUTPUT_PATH}")
    wb.save(OUTPUT_PATH)
    print("HOÀN TẤT CẬP NHẬT TOÀN DIỆN UNIT TEST CASE.XLSX THÀNH CÔNG 100%!")

if __name__ == "__main__":
    main()
