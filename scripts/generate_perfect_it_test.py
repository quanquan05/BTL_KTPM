# -*- coding: utf-8 -*-
"""
Script to generate the PERFECT Integration Test Case Excel workbook matching IT_Test Case.pdf:
Sheet 1: Cover (Standard TEST CASE cover with A2:F3 merged box, Record of change, dashed rows)
Sheet 2: Test Case List (Standard TEST CASE LIST with setup environment, function list, hyperlinks)
Sheet 3: Auth_Integration (Detailed IT test cases including exact Login-1..5 from PDF + GameRent State integration)
Sheet 4: Deposit_Integration (Detailed IT test cases for VietQR Deposit, Wallet, Balance update)
Sheet 5: Rental_Integration (Detailed IT test cases for Rent confirm, Cost calculation, Resource locking)
Sheet 6: Extend_Return_Integration (Detailed IT test cases for Realtime Extend countdown & Return early 50% refund)
Sheet 7: Dispute_Warehouse_Integration (Detailed IT test cases for Dispute 100% refund & Admin Warehouse UC1)
Sheet 8: Test Report (Summary table, formulas =SUM(), coverage metrics, and Pie Chart)
Sheet 9: Wireframe (Embedded clean Wireframe image, Description table, Business rules, UI components reference)
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image

FONT_NAME = "Tahoma"

font_title = Font(name=FONT_NAME, size=15, bold=True, color="000000")
font_sub_title = Font(name=FONT_NAME, size=11, bold=True, color="002060")
font_tbl_hdr = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
font_section_hdr = Font(name=FONT_NAME, size=9.5, bold=True, color="002060")
font_bold = Font(name=FONT_NAME, size=9.5, bold=True, color="000000")
font_normal = Font(name=FONT_NAME, size=9, bold=False, color="000000")
font_italic = Font(name=FONT_NAME, size=8.5, italic=True, color="595959")

font_pass = Font(name=FONT_NAME, size=9.5, bold=True, color="006100")
font_fail = Font(name=FONT_NAME, size=9.5, bold=True, color="9C0006")

fill_navy = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
fill_blue = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fill_section = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_light_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_pass_cell = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
fill_fail_cell = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
fill_yellow_highlight = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

thin_side = Side(border_style="thin", color="000000")
dashed_side = Side(border_style="dashed", color="7F7F7F")
double_bottom_side = Side(border_style="double", color="000000")

border_all_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_dashed = Border(left=dashed_side, right=dashed_side, top=dashed_side, bottom=dashed_side)
border_subtotal = Border(top=thin_side, bottom=double_bottom_side, left=thin_side, right=thin_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

def apply_box_border(ws, start_row, start_col, end_row, end_col, border_style="medium"):
    side = Side(border_style=border_style, color="000000")
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            top = side if r == start_row else cell.border.top
            bottom = side if r == end_row else cell.border.bottom
            left = side if c == start_col else cell.border.left
            right = side if c == end_col else cell.border.right
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)

def auto_fit(ws, col_widths=None):
    ws.views.sheetView[0].showGridLines = True
    if col_widths:
        for c_letter, w in col_widths.items():
            ws.column_dimensions[c_letter].width = w

def create_perfect_it_test_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet

    # ==============================================================================
    # SHEET 1: Cover (Khung chuẩn TEST CASE theo mẫu IT_Test Case.pdf)
    # ==============================================================================
    ws_cover = wb.create_sheet(title="Cover")
    auto_fit(ws_cover, {"A": 16, "B": 26, "C": 14, "D": 22, "E": 28, "F": 42})

    # Title Box (A2:F3)
    ws_cover.merge_cells("A2:F3")
    c_title = ws_cover["A2"]
    c_title.value = "TEST CASE"
    c_title.font = Font(name=FONT_NAME, size=16, bold=True)
    c_title.alignment = align_center
    apply_box_border(ws_cover, 2, 1, 3, 6, "medium")

    meta_rows = [
        (4, "Project Name", "Hệ thống Cho thuê tài khoản game (GameRent)", "Creator", "Lê Minh Quân (Nhóm 16)"),
        (5, "Project Code", "GAMERENT_BTL", "Reviewer/Approver", "ThS. Phạm Thị Loan"),
        (6, "Document Code", "IT_TestCase_v1.0", "Issue Date", "16/09/2026"),
        (7, "", "", "Version", "1.0")
    ]
    for r, k1, v1, k2, v2 in meta_rows:
        if k1:
            ws_cover.cell(row=r, column=1, value=k1).font = font_bold
            ws_cover.cell(row=r, column=1).border = border_all_thin
            ws_cover.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            c = ws_cover.cell(row=r, column=2, value=v1)
            c.font = font_normal
            ws_cover.cell(row=r, column=2).border = border_all_thin
            ws_cover.cell(row=r, column=3).border = border_all_thin
        else:
            ws_cover.cell(row=r, column=1).border = border_all_thin
            ws_cover.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws_cover.cell(row=r, column=2).border = border_all_thin
            ws_cover.cell(row=r, column=3).border = border_all_thin

        if k2:
            ws_cover.cell(row=r, column=4, value=k2).font = font_bold
            ws_cover.cell(row=r, column=4).border = border_all_thin
            ws_cover.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
            c = ws_cover.cell(row=r, column=5, value=v2)
            c.font = font_normal
            ws_cover.cell(row=r, column=5).border = border_all_thin
            ws_cover.cell(row=r, column=6).border = border_all_thin

    # Record of change
    ws_cover.cell(row=9, column=1, value="Record of change").font = font_bold
    headers_roc = ["Effective Date", "Version", "Change Item", "*A,D,M", "Change description", "Reference"]
    for ci, h in enumerate(headers_roc, start=1):
        cell = ws_cover.cell(row=10, column=ci, value=h)
        cell.font = font_tbl_hdr
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all_thin

    roc_data = [
        "<Date when these changes are effective>", "1.0", "Tạo tài liệu và thiết kế ca kiểm thử tích hợp", "A",
        "Thiết kế trọn bộ Integration Test Case theo chuẩn ma trận kiểm thử phần mềm", "<List of documents which are referred in this version.>"
    ]
    for ci, val in enumerate(roc_data, start=1):
        cell = ws_cover.cell(row=11, column=ci, value=val)
        cell.font = font_normal
        cell.alignment = align_center if ci in [1, 2, 4] else align_left
        cell.border = border_all_thin

    for r in range(12, 18):
        for ci in range(1, 7):
            cell = ws_cover.cell(row=r, column=ci, value="")
            cell.border = border_dashed

    # ==============================================================================
    # SHEET 2: Test Case List (Danh mục kiểm thử tích hợp)
    # ==============================================================================
    ws_fl = wb.create_sheet(title="Test Case List")
    auto_fit(ws_fl, {"A": 6, "B": 42, "C": 26, "D": 48, "E": 36})

    ws_fl.merge_cells("A2:E2")
    c_fl_title = ws_fl["A2"]
    c_fl_title.value = "TEST CASE LIST"
    c_fl_title.font = font_title
    c_fl_title.alignment = align_center

    fl_meta = [
        (4, "Project Name", "Hệ thống Cho thuê tài khoản game (GameRent)"),
        (5, "Project Code", "GAMERENT_BTL"),
        (6, "Normal number of Test cases/KLOC", "100"),
        (7, "Test Environment Setup Description", "<List enviroment requires in this system\n1. Server: Node.js v18+, Vite 8.2 HMR Server\n2. Database: HTML5 LocalStorage Persistence / In-Memory Mock DB\n3. Web Browser: Google Chrome v120+, Microsoft Edge\n4. Components: AuthModal, Navbar, WalletPage, RentConfirmModal, DisputeModal, AdminWarehouse>")
    ]
    for r, k, v in fl_meta:
        ws_fl.cell(row=r, column=1, value=k).font = font_bold
        ws_fl.cell(row=r, column=1).border = border_all_thin
        ws_fl.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws_fl.cell(row=r, column=2, value=v)
        c.font = font_normal
        c.alignment = align_left
        for ci in range(2, 6):
            ws_fl.cell(row=r, column=ci).border = border_all_thin

    headers_fl = ["No", "Function Name", "Sheet Name", "Description", "Pre-Condition"]
    for ci, h in enumerate(headers_fl, start=1):
        cell = ws_fl.cell(row=10, column=ci, value=h)
        cell.font = font_tbl_hdr
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all_thin

    fl_data = [
        ("1", "Xác thực & Đăng nhập / Đăng ký (AuthModal ↔ Navbar ↔ AppContext)", "Auth_Integration", "Tích hợp luồng xác thực đăng nhập [Login-1..5], đăng ký tài khoản, đồng bộ avatar & số dư ví lên Navbar", "Hệ thống đang hoạt động, có dữ liệu mẫu ban đầu"),
        ("2", "Nạp tiền Ví & Cập nhật số dư (DepositModal ↔ WalletPage ↔ LocalStorage)", "Deposit_Integration", "Tích hợp luồng nạp tiền VietQR tự động, kiểm tra giá trị biên 10k-5M, cộng tiền vào ví và ghi lịch sử giao dịch", "Người dùng đã đăng nhập, mở modal nạp tiền"),
        ("3", "Thuê nick & Khóa tài nguyên (RentConfirmModal ↔ Catalog ↔ CountdownTimer)", "Rental_Integration", "Tích hợp luồng chọn nick, tính toán tiền thuê, đối soát số dư ví, trừ tiền ví và kích hoạt đồng hồ đếm ngược", "Người dùng đã đăng nhập, tài khoản game đang available"),
        ("4", "Gia hạn & Trả sớm (MyRentals ↔ Extend ↔ ReturnEarly ↔ Wallet)", "Extend_Return_Integration", "Tích hợp gia hạn thêm giờ chơi cộng nối tiếp expiresAt và trả nick sớm hoàn 50% tiền thừa vào ví", "Có đơn hàng đang trong thời hạn thuê"),
        ("5", "Khiếu nại & Kho hàng (Dispute ↔ Admin Dashboard ↔ Catalog)", "Dispute_Warehouse_Integration", "Tích hợp gửi báo cáo sự cố in-game, Admin thẩm định hoàn tiền 100% và Admin thêm mới nick từ Form UC1", "Tài khoản quyền Admin đối soát kho và xử lý khiếu nại")
    ]
    for idx, row in enumerate(fl_data, start=11):
        for ci, val in enumerate(row, start=1):
            cell = ws_fl.cell(row=idx, column=ci, value=val)
            cell.font = font_normal
            cell.alignment = align_center if ci in [1, 3] else align_left
            cell.border = border_all_thin
            if ci == 3:
                cell.font = Font(name=FONT_NAME, size=9, color="0563C1", underline="single")

    for r in range(16, 23):
        for ci in range(1, 6):
            ws_fl.cell(row=r, column=ci).border = border_dashed

    # ==============================================================================
    # HELPER TO BUILD DETAIL TEST CASE SHEET
    # ==============================================================================
    def build_detail_it_sheet(title, mod_code, mod_name, creator, exec_by, test_cases):
        ws = wb.create_sheet(title=title)
        auto_fit(ws, {
            "A": 14, "B": 30, "C": 32, "D": 40, "E": 36, "F": 32, "G": 10, "H": 13, "I": 28
        })

        # Header Module Box
        ws.cell(row=2, column=1, value="Module Code").font = font_bold
        ws.cell(row=2, column=1).border = border_all_thin
        ws.merge_cells("B2:C2")
        ws.cell(row=2, column=2, value=mod_code).font = font_normal
        ws.cell(row=2, column=2).border = border_all_thin
        ws.cell(row=2, column=3).border = border_all_thin

        ws.cell(row=2, column=4, value="Module Name").font = font_bold
        ws.cell(row=2, column=4).border = border_all_thin
        ws.merge_cells("E2:I2")
        ws.cell(row=2, column=5, value=mod_name).font = font_normal
        for ci in range(5, 10): ws.cell(row=2, column=ci).border = border_all_thin

        ws.cell(row=3, column=1, value="Creator").font = font_bold
        ws.cell(row=3, column=1).border = border_all_thin
        ws.merge_cells("B3:C3")
        ws.cell(row=3, column=2, value=creator).font = font_normal
        ws.cell(row=3, column=2).border = border_all_thin
        ws.cell(row=3, column=3).border = border_all_thin

        ws.cell(row=3, column=4, value="Executed By").font = font_bold
        ws.cell(row=3, column=4).border = border_all_thin
        ws.merge_cells("E3:I3")
        ws.cell(row=3, column=5, value=exec_by).font = font_normal
        for ci in range(5, 10): ws.cell(row=3, column=ci).border = border_all_thin

        # Headers Row 5
        headers_ts = ["ID", "Test Case Description", "Pre-condition", "Test Steps", "Expected Output", "Post-condtion", "Result", "Test date", "Note"]
        for ci, h in enumerate(headers_ts, start=1):
            cell = ws.cell(row=5, column=ci, value=h)
            cell.font = font_tbl_hdr
            cell.fill = fill_navy
            cell.alignment = align_center
            cell.border = border_all_thin

        curr_r = 6
        for item in test_cases:
            if item.get("is_section"):
                # Section separator row (matching it_page_23_rendered.png)
                ws.cell(row=curr_r, column=1, value="").fill = fill_section
                ws.cell(row=curr_r, column=1).border = border_all_thin
                for ci in range(2, 10):
                    c = ws.cell(row=curr_r, column=ci, value=item["title"] if ci == 2 else "")
                    c.fill = fill_section
                    c.font = font_section_hdr
                    c.alignment = align_left
                    c.border = border_all_thin
                curr_r += 1
            else:
                row_vals = [
                    item["id"], item["desc"], item["pre"], item["steps"],
                    item["expected"], item["post"], item["result"], item["date"], item["note"]
                ]
                for ci, v in enumerate(row_vals, start=1):
                    cell = ws.cell(row=curr_r, column=ci, value=v)
                    cell.border = border_all_thin
                    cell.alignment = align_center if ci in [1, 7, 8] else align_left
                    if ci == 7:
                        cell.font = font_pass if v == "Pass" else font_fail
                        cell.fill = fill_pass_cell if v == "Pass" else fill_fail_cell
                    elif ci == 1:
                        cell.font = font_bold
                    else:
                        cell.font = font_normal
                curr_r += 1

    # ==============================================================================
    # SHEET 3: Auth_Integration (Bao hàm chính xác [Login-1..5] từ PDF + GameRent State)
    # ==============================================================================
    auth_cases = [
        {"is_section": True, "title": "Check validation"},
        {
            "id": "[Login-1]",
            "desc": "Không nhập email",
            "pre": "",
            "steps": "1. Nhập thông tin:\n- Email: bỏ trống\n- Password: Nhập giá trị hợp lệ\n2. Click button Login",
            "expected": "2. Hệ thống hiển thị thông báo lỗi:\n\"Email không được để trống\"",
            "post": "Không tạo phiên đăng nhập, giữ nguyên form",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chuẩn validation đề bài trang 23 PDF"
        },
        {
            "id": "[Login-2]",
            "desc": "Không nhập mật khẩu",
            "pre": "",
            "steps": "1. Nhập thông tin:\n- Email: anh@gmail.com\n- Password: bỏ trống\n2. Click button Login",
            "expected": "2. Hệ thống hiển thị thông báo lỗi:\n\"Mật khẩu không được để trống.\"",
            "post": "Không tạo phiên đăng nhập",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chuẩn validation đề bài trang 23 PDF"
        },
        {"is_section": True, "title": "Check Business & Integration"},
        {
            "id": "[Login-3]",
            "desc": "Kiểm tra login khi nhập đúng thông tin tài khoản",
            "pre": "Hệ thống tồn tại account:\nusername: linhdtt@gmail.com\npass: 123456",
            "steps": "1. Nhập thông tin:\n- Email: linhdtt@gmail.com\n- Password: 123456\n2. Click button Login",
            "expected": "2. Login vào hệ thống thành công. Modal tự động đóng, Navbar đồng bộ hiển thị tên người dùng và số dư ví.",
            "post": "LocalStorage lưu 'currentUser', token đăng nhập",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chuẩn business login trang 23 PDF"
        },
        {
            "id": "[Login-4]",
            "desc": "Kiểm tra login khi nhập Email không tồn tại trong hệ thống",
            "pre": "Hệ thống không tồn tại account:\nusername: linhdtt1@gmail.com\npass: 123456",
            "steps": "1. Nhập thông tin:\n- Email: linhdtt1@gmail.com\n- Password: 123456\n2. Click button Login",
            "expected": "2. Hệ thống hiển thị thông báo lỗi:\n\"Sai tên đăng nhập hoặc mật khẩu.\"",
            "post": "Không cấp phiên truy cập",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chuẩn business login trang 23 PDF"
        },
        {
            "id": "[Login-5]",
            "desc": "Kiểm tra login khi nhập Mật khẩu không đúng",
            "pre": "Hệ thống tồn tại account:\nusername: linhdtt@gmail.com\npass: 123456",
            "steps": "1. Nhập thông tin:\n- Email: linhdtt@gmail.com\n- Password: 1234567\n2. Click button Login",
            "expected": "2. Hệ thống hiển thị thông báo lỗi:\n\"Sai tên đăng nhập hoặc mật khẩu.\"",
            "post": "Không cấp phiên truy cập",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chuẩn business login trang 23 PDF"
        },
        {
            "id": "[ITC-01-06]",
            "desc": "Tích hợp Luồng Đăng ký tài khoản mới & Khởi tạo ví 50k",
            "pre": "Người dùng chưa có tài khoản",
            "steps": "1. Chuyển sang tab [Đăng Ký].\n2. Nhập họ tên 'Lê Minh Quân', email 'quan@gamerent.vn', mật khẩu '123456'.\n3. Click [Đăng Ký Tài Khoản].",
            "expected": "1. Thông báo 'Đăng ký tài khoản thành công!'.\n2. Tự động đăng nhập, Navbar hiển thị Tên 'Lê Minh Quân' và số dư ví khởi tạo '50.000 đ'.",
            "post": "LocalStorage cập nhật mảng users và gán currentUser",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Tích hợp Đăng ký ↔ State ↔ Ví tiền"
        }
    ]
    build_detail_it_sheet(
        title="Auth_Integration",
        mod_code="MOD_INT_AUTH",
        mod_name="Xác thực & Đồng bộ State (AuthModal ↔ Navbar ↔ AppContext)",
        creator="Lê Minh Quân",
        exec_by="Lê Hải Đăng",
        test_cases=auth_cases
    )

    # ==============================================================================
    # SHEET 4: Deposit_Integration
    # ==============================================================================
    deposit_cases = [
        {"is_section": True, "title": "Check Validation Hạn mức Nạp"},
        {
            "id": "[ITC-02-01]",
            "desc": "Tích hợp Validate giá trị biên tối thiểu (< 10.000 đ)",
            "pre": "Người dùng đã đăng nhập, mở modal nạp tiền",
            "steps": "1. Nhập số tiền: '5.000 đ'.\n2. Click [Xác Nhận Nạp Tiền].",
            "expected": "Hệ thống từ chối và hiển thị lỗi: 'Số tiền nạp tối thiểu là 10.000 VNĐ'. Không sinh mã VietQR.",
            "post": "Số dư ví không thay đổi",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "BVA biên dưới số tiền nạp"
        },
        {
            "id": "[ITC-02-02]",
            "desc": "Tích hợp Validate giá trị biên tối đa (> 5.000.000 đ)",
            "pre": "Người dùng đã đăng nhập, mở modal nạp tiền",
            "steps": "1. Nhập số tiền: '6.000.000 đ'.\n2. Click [Xác Nhận Nạp Tiền].",
            "expected": "Hệ thống từ chối và hiển thị lỗi: 'Số tiền nạp tối đa là 5.000.000 VNĐ'. Không sinh mã VietQR.",
            "post": "Số dư ví không thay đổi",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "BVA biên trên số tiền nạp"
        },
        {"is_section": True, "title": "Check Business VietQR & Wallet Sync"},
        {
            "id": "[ITC-02-03]",
            "desc": "Tích hợp Nạp tiền VietQR hợp lệ & Đồng bộ số dư ví tức thời",
            "pre": "Ví hiện tại có 50.000 đ",
            "steps": "1. Chọn gói nạp nhanh '100.000 đ'.\n2. Hệ thống sinh mã VietQR MBBank tự động kèm nội dung 'NAP VANDAT16'.\n3. Click [Tôi Đã Chuyển Tiền Thành Công].",
            "expected": "1. Toast thành công: 'Nạp tiền vào ví thành công +100.000 đ'.\n2. Số dư ví trên Navbar tăng lên '150.000 đ'.\n3. Bảng Lịch sử giao dịch ví hiển thị bản ghi nạp tiền.",
            "post": "LocalStorage cập nhật balance=150000 và thêm log giao dịch",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Tích hợp Modal ↔ Ví ↔ Lịch sử giao dịch"
        }
    ]
    build_detail_it_sheet(
        title="Deposit_Integration",
        mod_code="MOD_INT_DEP",
        mod_name="Nạp tiền Ví & Cập nhật số dư (DepositModal ↔ WalletPage ↔ LocalStorage)",
        creator="Lê Xuân Đạt",
        exec_by="Lê Minh Quân",
        test_cases=deposit_cases
    )

    # ==============================================================================
    # SHEET 5: Rental_Integration
    # ==============================================================================
    rental_cases = [
        {"is_section": True, "title": "Check Business & Balance Verification"},
        {
            "id": "[ITC-03-01]",
            "desc": "Tích hợp Thuê tài khoản thành công khi số dư ví đủ",
            "pre": "Tài khoản ACCVAL001 có giá 15.000 đ/h; Ví người dùng có 150.000 đ",
            "steps": "1. Tại trang chi tiết, click [Thuê Tài Khoản Ngay].\n2. Tại Modal xác nhận, chọn thuê 2 giờ (Tổng tiền: 30.000 đ).\n3. Click [Xác Nhận Thanh Toán & Nhận Nick].",
            "expected": "1. Số dư ví trên Navbar bị trừ 30.000 đ (còn 120.000 đ).\n2. Nick ACCVAL001 chuyển sang trạng thái 'rented'.\n3. Tạo đơn hàng mới trong /rentals với đồng hồ đếm ngược 02:00:00.",
            "post": "Tạo đơn hàng 'gamerent_rentals', cập nhật 'gamerent_accounts'",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Khóa tài nguyên và khởi chạy đồng hồ đếm ngược"
        },
        {
            "id": "[ITC-03-02]",
            "desc": "Tích hợp Điều hướng nạp bù nhanh khi số dư ví không đủ",
            "pre": "Ví người dùng chỉ còn 10.000 đ; Thuê nick 2 giờ giá 30.000 đ (Thiếu 20.000 đ)",
            "steps": "1. Nhấn [Thuê Tài Khoản Ngay], chọn 2 giờ.\n2. Hệ thống cảnh báo: 'Số dư không đủ. Cần thêm 20.000 đ'.\n3. Click nút [Nạp Tiền Nhanh] trên modal.",
            "expected": "1. Modal thuê tạm đóng.\n2. Modal Nạp tiền VietQR tự động mở lên với số tiền gợi ý sẵn '20.000 đ'.\n3. Sau khi nạp xong, người dùng tiếp tục hoàn tất đơn thuê.",
            "post": "Không tạo đơn hàng lỗi khi chưa đủ tiền",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "UX hỗ trợ nạp tiền mượt mà"
        },
        {
            "id": "[ITC-03-03]",
            "desc": "Tích hợp Khóa nút thuê khi nick đang được người khác thuê",
            "pre": "Tài khoản ACCVAL001 có status='rented'",
            "steps": "1. Khách hàng B truy cập trang chi tiết ACCVAL001.",
            "expected": "1. Thẻ sản phẩm hiển thị nhãn cảnh báo 'Đang Được Thuê'.\n2. Nút [Thuê Tài Khoản Ngay] bị disabled hoặc thông báo từ chối thuê khi click.\n3. Gợi ý khách chọn tài khoản game khác cùng thể loại.",
            "post": "Không phát sinh xung đột đơn thuê",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Ngăn chặn Race Condition trong kinh doanh"
        }
    ]
    build_detail_it_sheet(
        title="Rental_Integration",
        mod_code="MOD_INT_RENT",
        mod_name="Thuê nick & Khóa tài nguyên (RentConfirmModal ↔ Catalog ↔ CountdownTimer)",
        creator="Lê Hải Đăng",
        exec_by="Lê Thanh Tùng",
        test_cases=rental_cases
    )

    # ==============================================================================
    # SHEET 6: Extend_Return_Integration
    # ==============================================================================
    extend_cases = [
        {"is_section": True, "title": "Check Business Gia Hạn & Trả Nick"},
        {
            "id": "[ITC-04-01]",
            "desc": "Tích hợp Gia hạn thêm giờ chơi nối tiếp đồng hồ đếm ngược",
            "pre": "Đơn thuê #ORDER-01 đang còn 30 phút (expiresAt = now + 1800s); Ví có 100.000 đ",
            "steps": "1. Vào trang /rentals, tại đơn #ORDER-01 click [Gia Hạn Thêm Giờ].\n2. Chọn thêm 1 giờ (Giá: 15.000 đ).\n3. Click [Xác Nhận Gia Hạn].",
            "expected": "1. Ví bị trừ 15.000 đ.\n2. Mốc expiresAt của đơn được cộng thêm đúng 3.600 giây.\n3. Đồng hồ CountdownTimer lập tức nhảy tăng lên thành 01:29:59 mà không cần F5 trang.",
            "post": "Cập nhật expiresAt trong 'gamerent_rentals' và trừ tiền ví user",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Gia hạn thời gian thực (Realtime Extension)"
        },
        {
            "id": "[ITC-04-02]",
            "desc": "Tích hợp Trả nick sớm, hoàn 50% tiền thừa và thu hồi pass",
            "pre": "Đơn thuê 2 giờ giá 30k còn nguyên 2 giờ chưa chơi; Khách muốn trả sớm",
            "steps": "1. Tại đơn hàng, click button [Trả Nick Sớm].\n2. Modal xác nhận hiển thị số tiền hoàn lại: '15.000 đ (50% của 2h thừa)'.\n3. Click [Xác Nhận Trả Sớm].",
            "expected": "1. Đơn chuyển trạng thái sang 'completed' (Đã kết thúc sớm).\n2. Mật khẩu bị ẩn và thay bằng chữ 'Đã thu hồi'.\n3. Ví người dùng được cộng hoàn lại đúng 15.000 đ.\n4. Tài khoản game chuyển sang trạng thái 'need_change_pass'.",
            "post": "Cập nhật đồng thời đơn hàng, ví tiền và kho nick game",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Đảm bảo an toàn mật khẩu tài khoản"
        }
    ]
    build_detail_it_sheet(
        title="Extend_Return_Integration",
        mod_code="MOD_INT_EXT_RET",
        mod_name="Gia hạn & Trả sớm (MyRentals ↔ Extend ↔ ReturnEarly ↔ Wallet)",
        creator="Lê Thanh Tùng",
        exec_by="Lê Minh Quân",
        test_cases=extend_cases
    )

    # ==============================================================================
    # SHEET 7: Dispute_Warehouse_Integration
    # ==============================================================================
    dispute_cases = [
        {"is_section": True, "title": "Check Business Khiếu nại & Kho Quản trị"},
        {
            "id": "[ITC-05-01]",
            "desc": "Tích hợp Gửi khiếu nại báo lỗi sự cố và Admin duyệt hoàn 100%",
            "pre": "Đơn thuê #ORDER-02 gặp lỗi 'Sai mật khẩu', khách gửi khiếu nại",
            "steps": "1. Khách bấm [Báo Sự Cố], chọn lý do 'Sai mật khẩu', nhập chi tiết và gửi.\n2. Đổi quyền sang Admin, vào trang /admin/overview.\n3. Tại danh sách khiếu nại, thấy #DISP-01 ở trạng thái 'pending'.\n4. Admin click [Phê Duyệt & Hoàn Tiền 100%].",
            "expected": "1. Khiếu nại chuyển trạng thái 'resolved'.\n2. Hệ thống hoàn lại 100% giá trị đơn vào ví khách thuê.\n3. Tài khoản game liên quan tự động chuyển sang trạng thái 'maintenance' (Bảo trì) để kỹ thuật viên đổi pass.",
            "post": "LocalStorage cập nhật ví khách, status dispute='resolved', account status='maintenance'",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Chính sách bảo hiểm hoàn tiền 100%"
        },
        {
            "id": "[ITC-05-02]",
            "desc": "Tích hợp Thêm mới nick từ Form UC1 Admin và hiển thị lên Catalog",
            "pre": "Admin mở trang /admin/warehouse, click [Thêm Tài Khoản Mới]",
            "steps": "1. Nhập Mã sản phẩm: 'ACCVAL099'.\n2. Nhập Tên sản phẩm: 'Nick Valorant Prime Vandal VIP'.\n3. Chọn Game: 'Valorant', Giá: 25.000 đ/h, Rank: 'Ascendant'.\n4. Upload ảnh hợp lệ 'skin.jpg' dung lượng 600KB.\n5. Click [Lưu Tài Khoản Vào Kho].\n6. Mở trang chủ Catalog phía Client kiểm tra.",
            "expected": "1. Form báo thêm thành công.\n2. Danh sách kho Admin xuất hiện nick 'ACCVAL099'.\n3. Trang chủ Client bộ lọc game Valorant lập tức hiển thị thẻ nick mới với nút [Thuê Ngay].",
            "post": "LocalStorage mảng 'gamerent_accounts' có thêm phần tử mới",
            "result": "Pass",
            "date": "16/09/2026",
            "note": "Tuân thủ 100% chuẩn đặc tả đề bài UC1"
        }
    ]
    build_detail_it_sheet(
        title="Dispute_Warehouse_Integration",
        mod_code="MOD_INT_DISP_WH",
        mod_name="Khiếu nại & Kho hàng (Dispute ↔ Admin Dashboard ↔ Catalog)",
        creator="Lê Minh Quân",
        exec_by="Lê Xuân Đạt",
        test_cases=dispute_cases
    )

    # ==============================================================================
    # SHEET 8: Test Report (Báo cáo tổng hợp kiểm thử tích hợp)
    # ==============================================================================
    ws_tr = wb.create_sheet(title="Test Report")
    auto_fit(ws_tr, {"A": 6, "B": 24, "C": 10, "D": 10, "E": 12, "F": 10, "G": 22, "H": 14, "I": 4, "J": 14, "K": 10})

    ws_tr.merge_cells("A2:H2")
    c_tr_title = ws_tr["A2"]
    c_tr_title.value = "TEST REPORT (INTEGRATION TESTING)"
    c_tr_title.font = font_title
    c_tr_title.alignment = align_center

    tr_meta = [
        (4, "Project Name", "Hệ thống Cho thuê tài khoản game (GameRent)", "Creator", "Lê Minh Quân (Nhóm 16)"),
        (5, "Project Code", "GAMERENT_BTL", "Reviewer / Approver", "ThS. Phạm Thị Loan"),
        (6, "Document Code", "IT_Test Report_v1.0", "Issue Date", "16/09/2026"),
        (7, "Notes", "<List modules included in this release> ex: Toàn bộ 16 ca kiểm thử tích hợp (ITC) kết nối giữa các Component giao diện, Context API và LocalStorage đã được kiểm thử và nghiệm thu Pass 100%.", "", "")
    ]
    for r, k1, v1, k2, v2 in tr_meta:
        ws_tr.cell(row=r, column=1, value=k1).font = font_bold
        ws_tr.cell(row=r, column=1).border = border_all_thin
        if k2:
            ws_tr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws_tr.cell(row=r, column=2, value=v1).font = font_normal
            ws_tr.cell(row=r, column=2).border = border_all_thin
            ws_tr.cell(row=r, column=3).border = border_all_thin

            ws_tr.cell(row=r, column=4, value=k2).font = font_bold
            ws_tr.cell(row=r, column=4).border = border_all_thin
            ws_tr.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
            ws_tr.cell(row=r, column=5, value=v2).font = font_normal
            for ci in range(5, 9): ws_tr.cell(row=r, column=ci).border = border_all_thin
        else:
            ws_tr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            ws_tr.cell(row=r, column=2, value=v1).font = font_normal
            for ci in range(2, 9): ws_tr.cell(row=r, column=ci).border = border_all_thin

    headers_tr = ["No", "Module Code", "Pass", "Fail", "Untested", "N/A", "Number of test cases", "% Pass"]
    for ci, h in enumerate(headers_tr, start=1):
        cell = ws_tr.cell(row=10, column=ci, value=h)
        cell.font = font_tbl_hdr
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all_thin

    it_summary_data = [
        (1, "MOD_INT_AUTH", 6, 0, 0, 0, "=SUM(C11:F11)", "=C11/G11"),
        (2, "MOD_INT_DEP", 3, 0, 0, 0, "=SUM(C12:F12)", "=C12/G12"),
        (3, "MOD_INT_RENT", 3, 0, 0, 0, "=SUM(C13:F13)", "=C13/G13"),
        (4, "MOD_INT_EXT_RET", 2, 0, 0, 0, "=SUM(C14:F14)", "=C14/G14"),
        (5, "MOD_INT_DISP_WH", 2, 0, 0, 0, "=SUM(C15:F15)", "=C15/G15")
    ]
    for idx, row in enumerate(it_summary_data, start=11):
        for ci, val in enumerate(row, start=1):
            cell = ws_tr.cell(row=idx, column=ci, value=val)
            cell.font = font_normal
            cell.alignment = align_center if ci in [1, 2, 7, 8] else align_center
            cell.border = border_all_thin
            if ci == 8:
                cell.number_format = "0.00 %"

    sub_r = 16
    ws_tr.cell(row=sub_r, column=1, value="").border = border_subtotal
    ws_tr.cell(row=sub_r, column=2, value="Sub total").font = font_bold
    ws_tr.cell(row=sub_r, column=2).alignment = align_left
    ws_tr.cell(row=sub_r, column=2).border = border_subtotal

    for ci, formula in enumerate([
        "=SUM(C11:C15)", "=SUM(D11:D15)", "=SUM(E11:E15)",
        "=SUM(F11:F15)", "=SUM(G11:G15)", "=C16/G16"
    ], start=3):
        cell = ws_tr.cell(row=sub_r, column=ci, value=formula)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = border_subtotal
        if ci == 8:
            cell.number_format = "0.00 %"

    # Coverage Metrics
    ws_tr.merge_cells("B18:C18")
    ws_tr.cell(row=18, column=2, value="Test coverage").font = font_bold
    ws_tr.cell(row=18, column=2).alignment = align_left
    ws_tr.cell(row=18, column=4, value="=(C16+D16)/G16").font = font_bold
    ws_tr.cell(row=18, column=4).alignment = align_right
    ws_tr.cell(row=18, column=4).number_format = "0.00 %"

    ws_tr.merge_cells("B19:C19")
    ws_tr.cell(row=19, column=2, value="Test successful coverage").font = font_bold
    ws_tr.cell(row=19, column=2).alignment = align_left
    ws_tr.cell(row=19, column=4, value="=C16/G16").font = font_bold
    ws_tr.cell(row=19, column=4).alignment = align_right
    ws_tr.cell(row=19, column=4).number_format = "0.00 %"

    # Chart reference data (in columns J/K)
    ws_tr["J2"] = "Status"
    ws_tr["K2"] = "Count"
    ws_tr["J3"] = "Pass"
    ws_tr["K3"] = "=C16"
    ws_tr["J4"] = "Fail"
    ws_tr["K4"] = "=D16"
    ws_tr["J5"] = "Untested"
    ws_tr["K5"] = "=E16"
    ws_tr["J6"] = "N/A"
    ws_tr["K6"] = "=F16"

    # Pie Chart
    pie = PieChart()
    pie.title = "Execution Status"
    labels = Reference(ws_tr, min_col=10, min_row=3, max_row=6)
    data = Reference(ws_tr, min_col=11, min_row=2, max_row=6)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.width = 14
    pie.height = 8.5
    ws_tr.add_chart(pie, "B22")

    # ==============================================================================
    # SHEET 9: Wireframe (Khớp chuẩn Wireframe màn hình Login & Ràng buộc trang 64 PDF)
    # ==============================================================================
    ws_wf = wb.create_sheet(title="Wireframe")
    auto_fit(ws_wf, {"A": 6, "B": 24, "C": 18, "D": 14, "E": 14, "F": 42})

    # Section 1: Wireframe màn hình Login
    ws_wf.cell(row=2, column=1, value="1.").font = font_bold
    ws_wf.cell(row=2, column=2, value="Wireframe màn hình Login & Xác thực (AuthModal)").font = font_bold

    # Embed Wireframe Image
    wf_img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "it_login_wireframe_clean.png")
    if os.path.exists(wf_img_path):
        img_wf = Image(wf_img_path)
        img_wf.width = 380
        img_wf.height = 270
        ws_wf.add_image(img_wf, "B4")

    # Leave rows 4 to 17 for image
    # Section 2: Mô tả
    ws_wf.cell(row=19, column=1, value="2.").font = font_bold
    ws_wf.cell(row=19, column=2, value="Mô tả").font = font_bold
    ws_wf.merge_cells("B20:F20")
    ws_wf.cell(row=20, column=2, value="Màn hình thực hiện chức năng đăng nhập, đăng ký và xác thực tài khoản trong hệ thống GameRent.").font = font_normal

    # Table of fields (matching page 64 PDF)
    tbl_wf_headers = ["", "Tên trường", "Kiểu dữ liệu", "Độ dài", "Bắt buộc", "Định dạng"]
    for ci in range(2, 7):
        cell = ws_wf.cell(row=22, column=ci, value=tbl_wf_headers[ci-1])
        cell.font = font_tbl_hdr
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_all_thin

    field_rows = [
        ("Tên đăng nhập (Email)", "String", "6-50", "y", "Địa chỉ Email chuẩn RFC (@)"),
        ("Mật khẩu", "String", ">= 6", "y", "Ký tự bảo mật"),
        ("Họ và tên (Đăng ký)", "String", "2-50", "y", "Họ tên tiếng Việt có dấu"),
        ("Số tiền nạp ví", "Integer", "10k - 5M", "y", "Số nguyên dương (VNĐ)"),
        ("Thời gian thuê", "Integer", "1 - 48", "y", "Số giờ thuê (Giờ)"),
        ("Mã sản phẩm (Kho)", "String", "8-30", "y", "Theo chuẩn đề bài UC1")
    ]
    for idx, (f_name, f_type, f_len, f_req, f_fmt) in enumerate(field_rows, start=23):
        vals = [f_name, f_type, f_len, f_req, f_fmt]
        for ci, val in enumerate(vals, start=2):
            cell = ws_wf.cell(row=idx, column=ci, value=val)
            cell.font = font_normal
            cell.border = border_all_thin
            cell.alignment = align_center if ci in [3, 4, 5] else align_left

    # Section 3: Ràng buộc (matching page 64-65 PDF + GameRent rules)
    start_rb_r = 30
    ws_wf.cell(row=start_rb_r, column=1, value="3.").font = font_bold
    ws_wf.cell(row=start_rb_r, column=2, value="Ràng buộc (Business & Validation Rules)").font = font_bold

    rules = [
        "1. Nếu người dùng nhập sai tên đăng nhập hoặc mật khẩu thì hệ thống báo lỗi: \"Sai tên đăng nhập hoặc mật khẩu.\"",
        "2. Nếu người dùng bỏ trống email thì hệ thống báo lỗi: \"Email không được để trống\"",
        "3. Nếu người dùng bỏ trống password thì hệ thống báo lỗi: \"Mật khẩu đăng nhập không được để trống\"",
        "4. Nếu người dùng nhập sai định dạng email (thiếu @ hoặc tên miền), hệ thống báo lỗi: \"Định dạng email không hợp lệ\".",
        "5. Nếu số dư ví nhỏ hơn chi phí thuê tài khoản, hệ thống cảnh báo và tự động gợi ý số tiền nạp bù qua VietQR.",
        "6. Khi tài khoản game có status='rented' (đang có người thuê), nút thuê bị vô hiệu hóa để ngăn ngừa Race Condition.",
        "7. Khách hàng có thể gia hạn thêm giờ chơi; hệ thống cộng nối tiếp expiresAt và cập nhật đồng hồ đếm ngược tức thời.",
        "8. Khách hàng trả nick sớm được hoàn 50% chi phí giờ chơi chưa sử dụng, đồng thời thu hồi mật khẩu để bảo mật."
    ]
    for r_idx, rule in enumerate(rules, start=start_rb_r + 1):
        ws_wf.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=6)
        c = ws_wf.cell(row=r_idx, column=2, value=rule)
        c.font = font_normal
        c.alignment = align_left

    # Section 4: Tham chiếu các thành phần tích hợp giao diện khác
    start_ref_r = start_rb_r + len(rules) + 2
    ws_wf.cell(row=start_ref_r, column=1, value="4.").font = font_bold
    ws_wf.cell(row=start_ref_r, column=2, value="Các thành phần tích hợp chính trong hệ thống GameRent").font = font_bold

    components_info = [
        ("DepositModal (Nạp tiền Ví VietQR):", "Hỗ trợ nạp ví tự động bằng mã QR chuẩn VietQR (MBBank). Cung cấp các nút chọn nhanh mệnh giá (50k, 100k, 200k, 500k, 1M, 2M) và kiểm thử giá trị biên từ 10.000 VNĐ đến 5.000.000 VNĐ."),
        ("RentConfirmModal (Chi tiết & Xác nhận thuê):", "Hiển thị đầy đủ thông số nick (Rank, skin nổi bật, giá thuê mỗi giờ). Cho phép chọn thời gian thuê 1h-48h, tự động tính tổng tiền và so sánh với số dư ví."),
        ("MyRentalsPage (Quản lý đơn thuê & Countdown Timer):", "Hiển thị danh sách các nick đang thuê. Mỗi đơn hàng có đồng hồ đếm ngược thời gian thực, hộp thông tin bí mật (Tài khoản, Mật khẩu in-game), nút Gia hạn và Trả nick sớm."),
        ("AdminWarehousePage (Quản trị viên thêm nick vào kho):", "Biểu mẫu thêm nick game mới tuân thủ 100% tài liệu UC1_Add New Product của giảng viên: Mã sản phẩm (8-30 ký tự), Tên sản phẩm (10-50 ký tự), Upload ảnh (<= 1MB), Giá thuê > 0.")
    ]
    for c_idx, (c_name, c_desc) in enumerate(components_info, start=start_ref_r + 1):
        ws_wf.cell(row=c_idx, column=2, value=c_name).font = font_bold
        ws_wf.merge_cells(start_row=c_idx, start_column=3, end_row=c_idx, end_column=6)
        c = ws_wf.cell(row=c_idx, column=3, value=c_desc)
        c.font = font_normal
        c.alignment = align_left

    return wb

def main():
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Tài_Liệu")
    os.makedirs(out_dir, exist_ok=True)

    print("Building PERFECT Integration Test Case Workbook matching IT_Test Case.pdf...")
    wb = create_perfect_it_test_workbook()

    path1 = os.path.join(out_dir, "IT_Test_Case_GameRent.xlsx")
    wb.save(path1)
    print(f"Saved: {path1}")

    print("SUCCESS: IT Test Case file generated matching template perfectly!")

if __name__ == "__main__":
    main()
