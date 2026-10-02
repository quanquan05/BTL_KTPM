# -*- coding: utf-8 -*-
"""
Script chỉnh sửa và chuẩn hóa file IT_Test Case.xlsx theo mẫu của IT_Test Case Mẫu.xlsx
Đảm bảo 100% cấu trúc 5 sheet:
1. Cover
2. Test case List (Đầy đủ 10 phân hệ nghiệp vụ của dự án GameRent)
3. Login (12 Test cases chi tiết: Validation & Business logic)
4. Test Report (Công thức thống kê tự động kết nối Sheet Login)
5. Requrirement (Wireframe AuthModal thực tế + Đặc tả bảng trường & 7 Business Rules)

Dữ liệu đồng nhất hoàn toàn với dự án GameRent (Nhóm 15 - Lớp Kiểm thử phần mềm-1-1-26(N05)):
- Project Name: GameRent - Website Cho Thuê Tài Khoản Game Tự Động 24/7
- Project Code: GAMERENT
- Thành viên: Lê Minh Quân, Lê Hải Đăng, Lê Xuân Đạt, Lê Thanh Tùng
- Giảng viên hướng dẫn: ThS. Phạm Thị Loan
- Hình ảnh Wireframe AuthModal thực tế từ public/login_wireframe_real_sized.png
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

TEMPLATE_PATH = r"e:\BTL_KTPM\Tài_Liệu\IT_Test Case Mẫu.xlsx"
OUTPUT_PATH = r"e:\BTL_KTPM\Tài_Liệu\IT_Test Case.xlsx"
TEMP_OUTPUT_PATH = r"e:\BTL_KTPM\Tài_Liệu\IT_Test Case_updated.xlsx"
WIREFRAME_IMG_PATH = r"e:\BTL_KTPM\public\login_wireframe_real_sized.png"

def update_it_test_case():
    print(f"Đang đọc file mẫu: {TEMPLATE_PATH}")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    
    # Định nghĩa style chuẩn
    font_title = Font(name="Tahoma", size=20, bold=True, color="000000")
    font_label = Font(name="Tahoma", size=10, bold=True, color="993300")
    font_value_green = Font(name="Tahoma", size=10, bold=False, color="008000")
    font_value_black = Font(name="Tahoma", size=10, bold=False, color="000000")
    font_header_white = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    font_sub_bold = Font(name="Tahoma", size=10, bold=True, color="000080")
    font_stat_formula = Font(name="Tahoma", size=10, bold=True, color="0000FF")
    
    fill_navy = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
    fill_sub_header = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    fill_pass = PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D1D5DB")
    hair_border_side = Side(border_style="hair", color="CBD5E0")
    table_cell_border = Border(left=thin_border_side, right=thin_border_side, top=hair_border_side, bottom=hair_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # =========================================================================
    # 1. SHEET: Cover
    # =========================================================================
    print("Đang xử lý Sheet: Cover...")
    ws_cover = wb['Cover']
    
    ws_cover['B4'].value = "Project Name"
    ws_cover['C4'].value = "GameRent - Website Cho Thuê Tài Khoản Game Tự Động 24/7"
    ws_cover['F4'].value = "Creator"
    ws_cover['G4'].value = "Nhóm 15 (Lê Minh Quân, Lê Hải Đăng, Lê Xuân Đạt, Lê Thanh Tùng)"
    
    ws_cover['B5'].value = "Project Code"
    ws_cover['C5'].value = "GAMERENT"
    ws_cover['F5'].value = "Reviewer/Approver"
    ws_cover['G5'].value = "ThS. Phạm Thị Loan"
    
    ws_cover['B6'].value = "Document Code"
    ws_cover['C6'].value = '=C5&"_"&"ITC"&"_"&"v1.0"'
    ws_cover['F6'].value = "Issue Date"
    ws_cover['G6'].value = "26/09/2026"
    
    ws_cover['F7'].value = "Version"
    ws_cover['G7'].value = "1.0"
    
    ws_cover['B10'].value = "Record of change"
    ws_cover['B10'].font = font_label
    
    headers_cover = ["Effective Date", "Version", "Change Item", "*A,D,M", "Change description", "Reference"]
    for idx, h in enumerate(headers_cover, start=2):
        cell = ws_cover.cell(11, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    changes = [
        ("10/09/2026", "0.1", "Khởi tạo tài liệu và thiết kế ca kiểm thử", "A", 
         "Khởi tạo tài liệu kiểm thử cho phân hệ Đăng nhập (Login) và các luồng nghiệp vụ xác thực người dùng trong dự án GameRent theo biểu mẫu chuẩn.",
         "SRS GameRent v1.0, Tài liệu thiết kế hệ thống"),
        ("26/09/2026", "1.0", "Hoàn thiện Test Case, Test Report & Đặc tả Yêu cầu", "M",
         "Cập nhật toàn diện 12 ca kiểm thử (Validation & Business logic), tích hợp hình ảnh Wireframe AuthModal, thực thi kiểm thử 100% Pass và xuất báo cáo Test Report chuẩn mẫu.",
         "Mã nguồn AuthModal.jsx, Vitest auth.test.js, Báo cáo BTL KTPM")
    ]
    
    for r_idx, chg in enumerate(changes, start=12):
        for c_idx, val in enumerate(chg, start=2):
            cell = ws_cover.cell(r_idx, c_idx)
            cell.value = val
            cell.font = font_value_black
            cell.border = table_cell_border
            if c_idx in [2, 3, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
    for r in range(14, 30):
        for c in range(2, 8):
            cell = ws_cover.cell(r, c)
            if cell.value and str(cell.value).startswith("<"):
                cell.value = None

    ws_cover.column_dimensions['A'].width = 6
    ws_cover.column_dimensions['B'].width = 16
    ws_cover.column_dimensions['C'].width = 16
    ws_cover.column_dimensions['D'].width = 24
    ws_cover.column_dimensions['E'].width = 10
    ws_cover.column_dimensions['F'].width = 30
    ws_cover.column_dimensions['G'].width = 38

    # =========================================================================
    # 2. SHEET: Test case List (Đầy đủ 10 phân hệ nghiệp vụ GameRent)
    # =========================================================================
    print("Đang xử lý Sheet: Test case List...")
    ws_list = wb['Test case List']
    
    ws_list['D1'].value = "TEST CASE LIST"
    ws_list['D1'].font = font_title
    ws_list['D1'].alignment = Alignment(horizontal="center", vertical="center")
    
    ws_list['B3'].value = "Project Name"
    ws_list['D3'].value = "=Cover!C4"
    ws_list['D3'].font = font_value_green
    
    ws_list['B4'].value = "Project Code"
    ws_list['D4'].value = "=Cover!C5"
    ws_list['D4'].font = font_value_green
    
    ws_list['B5'].value = "Test Environment Setup Description"
    env_desc = (
        "Môi trường kiểm thử hệ thống & phân hệ GameRent:\n"
        "1. Nền tảng: Node.js v20.x, React 19.x, Vite Dev Server\n"
        "2. Cơ sở dữ liệu: LocalStorage Database Engine Mock (gamerent_users, gamerent_current_user)\n"
        "3. Trình duyệt: Google Chrome v120+, Microsoft Edge v120+\n"
        "4. Framework tự động: Vitest Automation Suite, React Testing Library, jsdom\n"
        "5. Màn hình kiểm thử: Desktop (1920x1080) và Responsive Web"
    )
    ws_list['D5'].value = env_desc
    ws_list['D5'].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws_list.row_dimensions[5].height = 80
    
    headers_list = ["No", "Function Name", "Sheet Name", "Description", "Pre-Condition"]
    for idx, h in enumerate(headers_list, start=2):
        cell = ws_list.cell(8, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    functions = [
        (1, "Đăng nhập hệ thống (Login)", "Login", 
         "Kiểm thử chức năng xác thực người dùng, kiểm tra validation email/mật khẩu, phân quyền Renter/Admin và xử lý chặn tài khoản vi phạm bị khóa (isBlocked).",
         "Hệ thống GameRent khởi chạy trên trình duyệt web, cơ sở dữ liệu tài khoản người dùng đã sẵn sàng trong LocalStorage."),
         
        (2, "Đăng ký tài khoản mới (Register)", "Login", 
         "Kiểm thử form đăng ký thành viên mới (validation họ tên, email, mật khẩu >= 6 ký tự) và luồng tự động đăng nhập cấp ví 50.000 VNĐ.",
         "Trình duyệt mở modal AuthModal tại tab Đăng ký, cơ sở dữ liệu chưa tồn tại email đăng ký."),
         
        (3, "Nạp tiền ví VietQR Napas247 (Deposit)", "IT_WALLET_FLOW", 
         "Kiểm thử tích hợp nạp tiền tự động qua cổng VietQR, sinh mã QR động theo hạn mức nạp và cộng dồn số dư ví thời gian thực.",
         "Khách hàng đã đăng nhập tài khoản cá nhân, mở modal Nạp tiền trên trang Ví điện tử."),
         
        (4, "Thuê tài khoản & Bàn giao pass 24/7 (Rental)", "IT_RENT_DELIVER", 
         "Kiểm thử luồng kiểm tra số dư ví, trừ tiền thuê, cập nhật trạng thái kho nick và bàn giao thông tin đăng nhập in-game tức thì.",
         "Số dư ví khách hàng đủ chi trả chi phí thuê, tài khoản game đang ở trạng thái Available."),
         
        (5, "Giám sát đếm ngược & Gia hạn giờ chơi (Timer)", "IT_TIMER_EXT", 
         "Kiểm thử đồng hồ đếm ngược CountdownTimer thời gian thực, hệ thống cảnh báo 3 cấp độ màu và tính năng gia hạn cộng dồn thời gian chơi.",
         "Đơn thuê đang ở trạng thái Active, phiên chơi game đang diễn ra."),
         
        (6, "Trả nick sớm & Khiếu nại bảo hiểm (Refund)", "IT_REFUND_DISPUTE", 
         "Kiểm thử chính sách hoàn tiền 50% thời gian chưa sử dụng khi trả nick sớm và quy trình Admin phê duyệt bồi hoàn 100% khi tài khoản gặp sự cố.",
         "Khách hàng đang có đơn thuê active hoặc gửi khiếu nại sự cố tài khoản trong 15 phút đầu."),
         
        (7, "Quản trị thêm nick vào kho (Form UC1)", "Product Management", 
         "Kiểm thử biểu mẫu Thêm mới tài khoản game chuẩn 100% tài liệu UC1_Add New Product (áp dụng BVA/EP kiểm thử giá trị biên và bảng quyết định).",
         "Người dùng đã đăng nhập với vai trò Quản trị viên (Admin), mở form thêm nick trong OverviewDashboard."),
         
        (8, "Quản lý khách hàng & Đồng bộ dữ liệu CRM", "Customer Management", 
         "Kiểm thử tra cứu, phân trang, lọc khách hàng, cập nhật trạng thái hoạt động và đảm bảo tính nhất quán dữ liệu giữa Client và Admin.",
         "Tài khoản Admin đã đăng nhập, danh mục khách hàng đã được nạp từ cơ sở dữ liệu."),
         
        (9, "Phân tách danh sách yêu thích độc lập (Favorites)", "IT_FAV_SEP", 
         "Kiểm thử cơ chế lưu trữ danh sách tài khoản game yêu thích phân tách độc lập theo định danh từng tài khoản trong LocalStorage.",
         "Nhiều tài khoản người dùng khác nhau cùng thực hiện thao tác lưu tài khoản yêu thích trên cùng một trình duyệt."),
         
        (10, "Đổi mật khẩu người dùng (Change Password)", "Settings", 
         "Kiểm thử chức năng đổi mật khẩu cá nhân, kiểm tra mật khẩu hiện tại, độ mạnh mật khẩu mới và cập nhật bảo mật an toàn.",
         "Người dùng đang đăng nhập hệ thống, mở trang Cài đặt (Settings).")
    ]
    
    for idx, fn in enumerate(functions, start=9):
        ws_list.cell(idx, 2).value = fn[0]
        ws_list.cell(idx, 2).alignment = Alignment(horizontal='center', vertical='center')
        ws_list.cell(idx, 2).font = font_value_black
        ws_list.cell(idx, 2).border = table_cell_border
        
        ws_list.cell(idx, 3).value = fn[1]
        ws_list.cell(idx, 3).alignment = Alignment(horizontal='left', vertical='center')
        ws_list.cell(idx, 3).font = font_value_black
        ws_list.cell(idx, 3).border = table_cell_border
        
        ws_list.cell(idx, 4).value = fn[2]
        ws_list.cell(idx, 4).alignment = Alignment(horizontal='center', vertical='center')
        ws_list.cell(idx, 4).font = font_value_black
        ws_list.cell(idx, 4).border = table_cell_border
        
        ws_list.cell(idx, 5).value = fn[3]
        ws_list.cell(idx, 5).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_list.cell(idx, 5).font = font_value_black
        ws_list.cell(idx, 5).border = table_cell_border
        
        ws_list.cell(idx, 6).value = fn[4]
        ws_list.cell(idx, 6).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_list.cell(idx, 6).font = font_value_black
        ws_list.cell(idx, 6).border = table_cell_border
        
        ws_list.row_dimensions[idx].height = 42

    # Xóa sạch các dòng rác phía sau (dòng 19 đến 40)
    for r in range(19, 40):
        for c in range(1, 10):
            cell = ws_list.cell(r, c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)
            
    ws_list.column_dimensions['B'].width = 8
    ws_list.column_dimensions['C'].width = 30
    ws_list.column_dimensions['D'].width = 22
    ws_list.column_dimensions['E'].width = 50
    ws_list.column_dimensions['F'].width = 45

    # =========================================================================
    # 3. SHEET: Login
    # =========================================================================
    print("Đang xử lý Sheet: Login...")
    ws_login = wb['Login']
    
    ws_login['A2'].value = "Module Code"
    ws_login['A2'].font = font_label
    ws_login['B2'].value = "Login"
    ws_login['B2'].font = font_value_black
    
    ws_login['A3'].value = "Test requirement"
    ws_login['A3'].font = font_label
    ws_login['B3'].value = "Kiểm thử toàn diện chức năng Đăng nhập (AuthModal): kiểm tra các ràng buộc dữ liệu đầu vào (Validation) và các luồng xử lý nghiệp vụ xác thực người dùng (Business Logic) theo yêu cầu đặc tả của dự án GameRent."
    ws_login['B3'].font = font_value_black
    ws_login['B3'].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_login.row_dimensions[3].height = 35
    
    ws_login['A4'].value = "Tester"
    ws_login['A4'].font = font_label
    ws_login['B4'].value = "Nhóm 15 (Lê Minh Quân, Lê Hải Đăng, Lê Xuân Đạt, Lê Thanh Tùng)"
    ws_login['B4'].font = font_value_black
    
    ws_login['A5'].value = "Pass"
    ws_login['B5'].value = "Fail"
    ws_login['C5'].value = "Untested"
    ws_login['D5'].value = "N/A"
    ws_login['E5'].value = "Number of Test cases"
    for c in range(1, 6):
        cell = ws_login.cell(5, c)
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    ws_login['A6'].value = '=COUNTIF(G10:G30,"Pass")'
    ws_login['B6'].value = '=COUNTIF(G10:G30,"Fail")'
    ws_login['C6'].value = '=E6-D6-B6-A6'
    ws_login['D6'].value = '=COUNTIF(G10:G30,"N/A")'
    ws_login['E6'].value = '=COUNTA(A10:A14)+COUNTA(A16:A22)'
    for c in range(1, 6):
        cell = ws_login.cell(6, c)
        cell.font = font_stat_formula
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_cell_border

    headers_tc = [
        "ID", "Test Case Description", "Pre-condition", "Test Steps",
        "Expected Output", "Post-condtion", "Result", "Test date", "Note"
    ]
    for idx, h in enumerate(headers_tc, start=1):
        cell = ws_login.cell(8, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    test_cases_validation = [
        ("Không nhập email",
         "Hệ thống mở trang chủ GameRent, modal Đăng nhập (AuthModal) đang hiển thị.",
         "1. Nhập thông tin:\n- Email: bỏ trống\n- Password: nhập '123456'\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống chặn gửi form.\n2. Hiển thị thông báo lỗi: \"Email không được để trống\".",
         "Form giữ nguyên trạng thái, con trỏ tự động focus vào trường Email.",
         "Pass", "26/09/2026", "Đã tự động hóa trong vitest auth.test.js"),
        
        ("Không nhập mật khẩu",
         "Modal Đăng nhập đang hiển thị.",
         "1. Nhập thông tin:\n- Email: 'hung.nguyen@gmail.com'\n- Password: bỏ trống\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống chặn gửi form.\n2. Hiển thị thông báo lỗi: \"Mật khẩu không được để trống\".",
         "Form giữ nguyên dữ liệu email đã nhập, con trỏ focus vào trường Password.",
         "Pass", "26/09/2026", "Đã tự động hóa trong vitest auth.test.js"),
        
        ("Bỏ trống cả email và mật khẩu",
         "Modal Đăng nhập đang hiển thị.",
         "1. Để trống cả trường Email và Password.\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống chặn gửi form.\n2. Hiển thị thông báo lỗi: \"Email không được để trống\".",
         "Form không submit, focus vào trường Email.",
         "Pass", "26/09/2026", "Đã tự động hóa trong vitest auth.test.js"),
        
        ("Nhập email không đúng định dạng",
         "Modal Đăng nhập đang hiển thị.",
         "1. Nhập thông tin:\n- Email: 'gamerent.user.vn' (thiếu @)\n- Password: 'password123'\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống chặn submit form.\n2. Trình duyệt hiển thị cảnh báo định dạng Email hợp lệ trên ô nhập liệu HTML5.",
         "Form giữ nguyên để người dùng sửa lại địa chỉ Email.",
         "Pass", "26/09/2026", "Validation định dạng RFC 5322"),
        
        ("Nhập mật khẩu dưới 6 ký tự",
         "Modal Đăng nhập đang hiển thị.",
         "1. Nhập thông tin:\n- Email: 'hung.nguyen@gmail.com'\n- Password: '123' (3 ký tự, dưới biên 6)\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống kiểm tra và báo lỗi: \"Mật khẩu không chính xác.\" hoặc yêu cầu mật khẩu tối thiểu 6 ký tự.",
         "Form không đăng nhập, giữ nguyên để nhập lại.",
         "Pass", "26/09/2026", "BVA kiểm thử giá trị biên")
    ]
    
    test_cases_business = [
        ("Kiểm tra đăng nhập khi nhập đúng tài khoản Khách hàng (Renter)",
         "Hệ thống tồn tại tài khoản khách hàng:\nemail: hung.nguyen@gmail.com\npassword: tester123\nisBlocked: false",
         "1. Nhập thông tin:\n- Email: hung.nguyen@gmail.com\n- Password: tester123\n2. Click button 'Đăng nhập'.",
         "1. Đăng nhập vào hệ thống thành công, đóng modal AuthModal.\n2. Header cập nhật hiển thị tên \"Nguyễn Văn Hùng\", avatar và số dư ví.\n3. Lưu phiên đăng nhập vào LocalStorage (gamerent_current_user).",
         "Người dùng duy trì trạng thái đăng nhập với vai trò 'renter'.",
         "Pass", "26/09/2026", "Normal case - Đã kiểm thử tự động"),
        
        ("Kiểm tra đăng nhập khi nhập đúng tài khoản Quản trị viên (Admin)",
         "Hệ thống tồn tại tài khoản quản trị:\nemail: admin@gamerent.vn\npassword: admin123\nrole: admin",
         "1. Nhập thông tin:\n- Email: admin@gamerent.vn\n- Password: admin123\n2. Click button 'Đăng nhập'.",
         "1. Đăng nhập thành công với vai trò Quản trị viên.\n2. Header hiển thị tên \"Quản Lý\" và nút truy cập Admin Dashboard.\n3. Cấp quyền truy cập quản lý kho nick, đơn thuê, duyệt khiếu nại.",
         "Người dùng có quyền hạn admin trong toàn bộ phiên làm việc.",
         "Pass", "26/09/2026", "Phân quyền RBAC chính xác"),
        
        ("Kiểm tra đăng nhập khi nhập Email không tồn tại trong hệ thống",
         "Hệ thống không tồn tại tài khoản có email: nonexistent@gamerent.vn",
         "1. Nhập thông tin:\n- Email: nonexistent@gamerent.vn\n- Password: password123\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống hiển thị thông báo lỗi: \"Email không tồn tại trên hệ thống.\".\n2. Không tạo phiên làm việc trong LocalStorage.",
         "Trạng thái người dùng vẫn là khách vãng lai (Guest).",
         "Pass", "26/09/2026", "Xác thực tài khoản không tồn tại"),
        
        ("Kiểm tra đăng nhập khi nhập Mật khẩu không đúng",
         "Hệ thống tồn tại tài khoản:\nemail: admin@gamerent.vn\npassword: admin123",
         "1. Nhập thông tin:\n- Email: admin@gamerent.vn\n- Password: wrongpassword999\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống hiển thị thông báo lỗi: \"Mật khẩu không chính xác.\".\n2. Mật khẩu sai bị từ chối, form giữ nguyên email.",
         "Người dùng không đăng nhập được.",
         "Pass", "26/09/2026", "Xác thực mật khẩu sai"),
        
        ("Kiểm tra đăng nhập với Tài khoản bị khóa do vi phạm quy chế",
         "Hệ thống tồn tại tài khoản bị khóa:\nemail: blocked@test.com\npassword: 123456\nisBlocked: true",
         "1. Nhập thông tin:\n- Email: blocked@test.com\n- Password: 123456\n2. Click button 'Đăng nhập'.",
         "1. Hệ thống chặn đăng nhập.\n2. Hiển thị thông báo lỗi: \"Tài khoản của bạn đang bị khóa do vi phạm quy chế.\".\n3. Không cấp token phiên làm việc.",
         "Tài khoản bị khóa bị từ chối đăng nhập hoàn toàn.",
         "Pass", "26/09/2026", "Bảo mật và kiểm soát tài khoản vi phạm"),
        
        ("Kiểm tra luồng Đăng ký tài khoản mới tự động chuyển sang Đăng nhập",
         "Người dùng chưa có tài khoản, đang mở tab Đăng ký trên AuthModal.",
         "1. Tại AuthModal, click chuyển sang tab 'Đăng ký'.\n2. Nhập họ tên 'Lê Minh Quân', email 'newbie2026@gamerent.vn', mật khẩu 'pass123456', xác nhận 'pass123456'.\n3. Click 'Đăng ký ngay'.",
         "1. Đăng ký thành công, hệ thống tự động đăng nhập luôn mà không bắt nhập lại.\n2. Số dư ví khởi tạo được cộng ngay 50.000 VNĐ.\n3. Header hiển thị tên tài khoản mới.",
         "Tài khoản mới lưu vào LocalStorage và duy trì phiên đăng nhập.",
         "Pass", "26/09/2026", "Tích hợp luồng Đăng ký ↔ Đăng nhập tự động"),
        
        ("Kiểm tra chức năng Đăng xuất (Logout) và xóa phiên làm việc",
         "Người dùng đang ở trạng thái đăng nhập với tài khoản 'hung.nguyen@gmail.com'.",
         "1. Tại Header/Navbar, mở menu người dùng.\n2. Click chọn 'Đăng xuất'.",
         "1. Hệ thống xóa key 'gamerent_current_user' trong LocalStorage.\n2. Header chuyển về nút 'Đăng nhập / Đăng ký'.\n3. Trả về trạng thái người dùng khách an toàn.",
         "Phiên làm việc kết thúc hoàn toàn.",
         "Pass", "26/09/2026", "Quản lý vòng đời phiên người dùng")
    ]
    
    ws_login['A9'].value = ""
    ws_login['B9'].value = "Check validation"
    ws_login['B9'].font = font_sub_bold
    ws_login['B9'].fill = fill_sub_header
    for c in range(1, 10):
        if c > 2:
            ws_login.cell(9, c).value = None
        ws_login.cell(9, c).border = table_cell_border
        
    cur_row = 10
    for idx, tc in enumerate(test_cases_validation, start=1):
        ws_login.cell(cur_row, 1).value = f"[Login-{idx:02d}]"
        ws_login.cell(cur_row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws_login.cell(cur_row, 1).font = font_value_black
        
        for c_idx, val in enumerate(tc, start=2):
            cell = ws_login.cell(cur_row, c_idx)
            cell.value = val
            cell.font = font_value_black
            if c_idx == 7: # Result
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.fill = fill_pass
                cell.font = Font(name="Tahoma", size=9.5, bold=True, color="22543D")
            elif c_idx == 8: # Date
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = table_cell_border
        cur_row += 1
        
    ws_login.cell(cur_row, 1).value = ""
    ws_login.cell(cur_row, 2).value = "Check Business"
    ws_login.cell(cur_row, 2).font = font_sub_bold
    ws_login.cell(cur_row, 2).fill = fill_sub_header
    for c in range(1, 10):
        if c > 2:
            ws_login.cell(cur_row, c).value = None
        ws_login.cell(cur_row, c).border = table_cell_border
    cur_row += 1
    
    for idx, tc in enumerate(test_cases_business, start=6):
        ws_login.cell(cur_row, 1).value = f"[Login-{idx:02d}]"
        ws_login.cell(cur_row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws_login.cell(cur_row, 1).font = font_value_black
        
        for c_idx, val in enumerate(tc, start=2):
            cell = ws_login.cell(cur_row, c_idx)
            cell.value = val
            cell.font = font_value_black
            if c_idx == 7: # Result
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.fill = fill_pass
                cell.font = Font(name="Tahoma", size=9.5, bold=True, color="22543D")
            elif c_idx == 8: # Date
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = table_cell_border
        cur_row += 1
        
    for r in range(cur_row, 100):
        for c in range(1, 15):
            cell = ws_login.cell(r, c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)
            
    ws_login.column_dimensions['A'].width = 12
    ws_login.column_dimensions['B'].width = 28
    ws_login.column_dimensions['C'].width = 28
    ws_login.column_dimensions['D'].width = 36
    ws_login.column_dimensions['E'].width = 38
    ws_login.column_dimensions['F'].width = 26
    ws_login.column_dimensions['G'].width = 10
    ws_login.column_dimensions['H'].width = 13
    ws_login.column_dimensions['I'].width = 26

    # =========================================================================
    # 4. SHEET: Test Report
    # =========================================================================
    print("Đang xử lý Sheet: Test Report...")
    ws_report = wb['Test Report']
    
    ws_report['B1'].value = "TEST REPORT"
    ws_report['B1'].font = font_title
    ws_report['B1'].alignment = Alignment(horizontal="center", vertical="center")
    
    ws_report['B3'].value = "Project Name"
    ws_report['C3'].value = "=Cover!C4"
    ws_report['C3'].font = font_value_green
    ws_report['E3'].value = "Creator"
    ws_report['G3'].value = "=Cover!G4"
    ws_report['G3'].font = font_value_black
    
    ws_report['B4'].value = "Project Code"
    ws_report['C4'].value = "=Cover!C5"
    ws_report['C4'].font = font_value_green
    ws_report['E4'].value = "Reviewer/Approver"
    ws_report['G4'].value = "=Cover!G5"
    ws_report['G4'].font = font_value_black
    
    ws_report['B5'].value = "Document Code"
    ws_report['C5'].value = '=C4&"_"&"Test Report"&"_"&"v1.0"'
    ws_report['C5'].font = font_value_green
    ws_report['E5'].value = "Issue Date"
    ws_report['G5'].value = "26/09/2026"
    ws_report['G5'].font = font_value_black
    ws_report['H5'].value = None
    
    ws_report['B6'].value = "Notes"
    ws_report['C6'].value = "Báo cáo tổng hợp kết quả thực thi kiểm thử chức năng Đăng nhập (Login Module) của hệ thống GameRent. Toàn bộ 12 test cases đã được thực thi và đạt tỷ lệ thành công 100% (Pass)."
    ws_report['C6'].font = font_value_black
    ws_report['C6'].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_report.row_dimensions[6].height = 35
    
    headers_report = ["No", "Module code", "Pass", "Fail", "Untested", "N/A", "Number of  test cases"]
    for idx, h in enumerate(headers_report, start=2):
        cell = ws_report.cell(10, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    ws_report['B11'].value = 1
    ws_report['B11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['C11'].value = "=Login!B2"
    ws_report['C11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['D11'].value = "=Login!A6"
    ws_report['D11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['E11'].value = "=Login!B6"
    ws_report['E11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['F11'].value = "=Login!C6"
    ws_report['F11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['G11'].value = "=Login!D6"
    ws_report['G11'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['H11'].value = "=Login!E6"
    ws_report['H11'].alignment = Alignment(horizontal="center", vertical="center")
    for c in range(2, 9):
        ws_report.cell(11, c).border = table_cell_border

    for r in range(12, 14):
        for c in range(2, 10):
            cell = ws_report.cell(r, c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)
            
    ws_report['C14'].value = "Sub total"
    ws_report['C14'].font = font_header_white
    ws_report['C14'].fill = fill_navy
    ws_report['C14'].alignment = Alignment(horizontal="center", vertical="center")
    
    ws_report['D14'].value = "=SUM(D11:D13)"
    ws_report['E14'].value = "=SUM(E11:E13)"
    ws_report['F14'].value = "=SUM(F11:F13)"
    ws_report['G14'].value = "=SUM(G11:G13)"
    ws_report['H14'].value = "=SUM(H11:H13)"
    for c in range(4, 9):
        cell = ws_report.cell(14, c)
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    ws_report['C16'].value = "Test coverage"
    ws_report['C16'].font = font_label
    ws_report['E16'].value = "=(D14+E14)*100/(H14-G14)"
    ws_report['E16'].font = font_stat_formula
    ws_report['E16'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['F16'].value = "%"
    
    ws_report['C17'].value = "Test successful coverage"
    ws_report['C17'].font = font_label
    ws_report['E17'].value = "=D14*100/(H14-G14)"
    ws_report['E17'].font = font_stat_formula
    ws_report['E17'].alignment = Alignment(horizontal="center", vertical="center")
    ws_report['F17'].value = "%"
    
    ws_report.column_dimensions['B'].width = 8
    ws_report.column_dimensions['C'].width = 20
    ws_report.column_dimensions['D'].width = 12
    ws_report.column_dimensions['E'].width = 12
    ws_report.column_dimensions['F'].width = 12
    ws_report.column_dimensions['G'].width = 12
    ws_report.column_dimensions['H'].width = 22

    # =========================================================================
    # 5. SHEET: Requrirement
    # =========================================================================
    print("Đang xử lý Sheet: Requrirement...")
    ws_req = wb['Requrirement']
    
    for r in range(20, 60):
        for c in range(1, 10):
            cell = ws_req.cell(r, c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)
            
    ws_req['A2'].value = "1."
    ws_req['A2'].font = Font(name="Tahoma", size=11, bold=True)
    ws_req['B2'].value = "Wireframe & Giao diện luồng nghiệp vụ Đăng nhập (AuthModal GameRent)"
    ws_req['B2'].font = Font(name="Tahoma", size=11, bold=True, color="000080")
    
    ws_req._images.clear()
    if os.path.exists(WIREFRAME_IMG_PATH):
        img = Image(WIREFRAME_IMG_PATH)
        img.width = 300
        img.height = 384
        ws_req.add_image(img, 'B4')
        print(f"  -> Đã nhúng ảnh wireframe thành công: {WIREFRAME_IMG_PATH}")
    
    for r in range(4, 21):
        ws_req.row_dimensions[r].height = 18

    ws_req['A22'].value = "2."
    ws_req['A22'].font = Font(name="Tahoma", size=11, bold=True)
    ws_req['B22'].value = "Mô tả màn hình Đăng nhập"
    ws_req['B22'].font = Font(name="Tahoma", size=11, bold=True, color="000080")
    
    ws_req['B23'].value = "Màn hình thực hiện chức năng xác thực người dùng (AuthModal), hỗ trợ đăng nhập tài khoản Khách hàng (Renter) và Quản trị viên (Admin), kiểm tra mật khẩu và quản lý phiên làm việc tập trung."
    ws_req['B23'].font = font_value_black
    ws_req.row_dimensions[23].height = 25
    
    headers_fields = ["Tên trường", "Kiểu dữ liệu", "Độ dài", "Bắt buộc", "Định dạng"]
    for idx, h in enumerate(headers_fields, start=2):
        cell = ws_req.cell(25, idx)
        cell.value = h
        cell.font = font_header_white
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        
    field_data = [
        ("Tên đăng nhập / Email", "String", "6 - 50", "y", "Địa chỉ Email hợp lệ (chứa @ và tên miền)"),
        ("Mật khẩu", "String", "6 - 32", "y", "Ký tự bảo mật (hỗ trợ chữ, số, ký tự đặc biệt)"),
        ("Ghi nhớ phiên đăng nhập", "Boolean", "1", "n", "Checkbox lưu LocalStorage (true/false)")
    ]
    
    for r_idx, f_row in enumerate(field_data, start=26):
        for c_idx, val in enumerate(f_row, start=2):
            cell = ws_req.cell(r_idx, c_idx)
            cell.value = val
            cell.font = font_value_black
            cell.border = table_cell_border
            if c_idx in [3, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    ws_req['A30'].value = "3."
    ws_req['A30'].font = Font(name="Tahoma", size=11, bold=True)
    ws_req['B30'].value = "Ràng buộc nghiệp vụ (Business Rules)"
    ws_req['B30'].font = Font(name="Tahoma", size=11, bold=True, color="000080")
    
    rules = [
        ("1.", "Nếu người dùng bỏ trống trường Email thì hệ thống hiển thị thông báo lỗi: \"Email không được để trống\"."),
        ("2.", "Nếu người dùng bỏ trống trường Mật khẩu thì hệ thống hiển thị thông báo lỗi: \"Mật khẩu không được để trống\"."),
        ("3.", "Nếu người dùng nhập sai mật khẩu đăng nhập thì hệ thống hiển thị thông báo lỗi: \"Mật khẩu không chính xác.\"."),
        ("4.", "Nếu người dùng nhập Email chưa đăng ký trên hệ thống thì báo lỗi: \"Email không tồn tại trên hệ thống.\"."),
        ("5.", "Nếu tài khoản người dùng có trạng thái bị khóa (isBlocked = true) do vi phạm quy chế thì hệ thống chặn truy cập và thông báo: \"Tài khoản của bạn đang bị khóa do vi phạm quy chế.\"."),
        ("6.", "Đăng nhập thành công sẽ lưu thông tin phiên làm việc vào LocalStorage (key: 'gamerent_current_user'), cập nhật Header/Navbar với tên người dùng, avatar, số dư ví và phân quyền tương ứng (Renter hoặc Admin)."),
        ("7.", "Khi người dùng nhấn 'Đăng xuất', hệ thống phải xóa hoàn toàn thông tin phiên trong LocalStorage và chuyển giao diện về trạng thái chưa đăng nhập an toàn.")
    ]
    
    for r_idx, (r_num, r_text) in enumerate(rules, start=31):
        cell_num = ws_req.cell(r_idx, 2)
        cell_num.value = r_num
        cell_num.font = font_value_black
        cell_num.alignment = Alignment(horizontal="center", vertical="top")
        
        cell_text = ws_req.cell(r_idx, 3)
        cell_text.value = r_text
        cell_text.font = font_value_black
        cell_text.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws_req.row_dimensions[r_idx].height = 24
        
    ws_req.column_dimensions['A'].width = 5
    ws_req.column_dimensions['B'].width = 26
    ws_req.column_dimensions['C'].width = 20
    ws_req.column_dimensions['D'].width = 14
    ws_req.column_dimensions['E'].width = 12
    ws_req.column_dimensions['F'].width = 46

    # Lưu file: thử lưu vào OUTPUT_PATH trước, nếu bị khóa do Excel đang mở thì lưu vào TEMP_OUTPUT_PATH
    try:
        wb.save(OUTPUT_PATH)
        print(f"Lưu trực tiếp thành công vào: {OUTPUT_PATH}")
    except PermissionError:
        print(f"CẢNH BÁO: File {OUTPUT_PATH} đang được mở trong Excel (bị khóa ghi).")
        wb.save(TEMP_OUTPUT_PATH)
        print(f"Đã lưu bản cập nhật hoàn chỉnh vào: {TEMP_OUTPUT_PATH}")
        return False
    return True

if __name__ == "__main__":
    success = update_it_test_case()
    if success:
        print("HOÀN TẤT 100%!")
    else:
        print("Vui lòng đóng file Excel để ghi đè!")
